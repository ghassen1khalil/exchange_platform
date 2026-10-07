# -*- coding: utf-8 -*-
"""Mode « diapos modeles » : une diapo d'exemple du template sert de patron.

On clone la diapo (decor, couleurs, polices, icones, images fixes...) et on remplace seulement le
texte / les images / les graphiques / les tableaux « a trous », en conservant la mise en forme
du texte d'exemple (police, taille, couleur, puces).
"""
from __future__ import annotations

import copy
import re

from lxml import etree
from pptx.oxml import parse_xml
from pptx.util import Emu

from .common import A_NS, NSMAP, P_NS, R_NS, EMU_PER_CM, order_reading, to_cm
from .icons import remap_rels, renumber_ids

_DECOR_NUM = re.compile(r"^\s*\d{1,2}\s*[.)]?\s*$")
_SKIP_PH = {"sldNum", "ftr", "dt", "hdr"}
_TITLE_PH = {"title", "ctrTitle"}
_QA = "{%s}" % A_NS
_QP = "{%s}" % P_NS
_QR = "{%s}" % R_NS
_CHART_URI = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_TABLE_URI = "http://schemas.openxmlformats.org/drawingml/2006/table"


class PSlot(object):
    def __init__(self, kind, shape, sample="", sizes=None, paras=None):
        self.kind = kind            # title | text | picture | chart | table
        self.shape = shape
        self.el = shape._element
        self.name = shape.name
        self.left = shape.left or 0
        self.top = shape.top or 0
        self.width = shape.width or 0
        self.height = shape.height or 0
        self.sample = sample
        self.paras = paras or []    # [(taille_pt|None, gras, niveau, texte)]
        self.ph_empty_picture = False

    @property
    def size_pt(self):
        sizes = [p[0] for p in self.paras if p[0]]
        return max(sizes) if sizes else None


def iter_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shp in shapes:
        yield shp
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in iter_shapes(shp.shapes):
                yield sub


def _ph_type(el):
    ph = el.find(".//" + _QP + "nvPr/" + _QP + "ph")
    if ph is None:
        return None
    return ph.get("type") or "obj"


def _para_info(shape):
    out = []
    for p in shape.text_frame.paragraphs:
        txt = "".join(r.text for r in p.runs).strip()
        if not txt:
            continue
        size = None
        bold = False
        for r in p.runs:
            if r.font.size:
                size = r.font.size.pt
            if r.font.bold:
                bold = True
            break
        out.append((size, bold, p.level or 0, txt))
    return out


def analyze_slide(slide, slide_w, slide_h):
    """Retourne la liste ordonnee des zones remplissables d'une diapo modele."""
    titles, texts, pics, charts, tables = [], [], [], [], []
    for shp in iter_shapes(slide.shapes):
        el = shp._element
        pht = _ph_type(el)
        if pht in _SKIP_PH:
            continue
        tag = etree.QName(el).localname
        if tag == "graphicFrame":
            uri = el.find(".//" + _QA + "graphicData")
            uri = uri.get("uri") if uri is not None else ""
            if uri == _CHART_URI:
                charts.append(PSlot("chart", shp))
            elif uri == _TABLE_URI:
                tables.append(PSlot("table", shp))
            continue
        if tag == "pic":
            if (shp.width or 0) * (shp.height or 0) > 0.06 * slide_w * slide_h:
                pics.append(PSlot("picture", shp))
            continue
        if tag == "grpSp":
            continue
        if pht in ("pic", "media") and not shp.has_text_frame:
            ps = PSlot("picture", shp)
            ps.ph_empty_picture = True
            pics.append(ps)
            continue
        if pht == "pic":
            ps = PSlot("picture", shp)
            ps.ph_empty_picture = True
            pics.append(ps)
            continue
        if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
            txt = shp.text_frame.text.strip()
            if not txt or _DECOR_NUM.match(txt):
                continue
            paras = _para_info(shp)
            ps = PSlot("text", shp, sample=txt.replace("\n", " / "), paras=paras)
            if pht in _TITLE_PH:
                ps.kind = "title"
                titles.append(ps)
            else:
                texts.append(ps)
    title = titles[0] if titles else None
    if title is None and texts:
        top = sorted(texts, key=lambda s: (s.top, s.left))[0]
        biggest = max([s.size_pt or 0 for s in texts] or [0])
        if (top.size_pt or 0) >= 0.9 * biggest and top.top < 0.35 * slide_h:
            top.kind = "title"
            title = top
            texts.remove(top)
    body = order_reading(texts + pics + charts + tables, slide_h)
    return title, body


def describe_slot(s):
    if s.kind in ("title", "text"):
        sizes = ", ".join("%s%s" % ((("%gpt" % p[0]) if p[0] else "?"), " gras" if p[1] else "") for p in s.paras[:4])
        n = len(s.sample)
        return "texte, %d paragraphe(s) d'exemple [%s], ~%d car. attendus (max ~%d)" % (len(s.paras), sizes, n, int(n * 1.3) + 5)
    if s.kind == "picture":
        return "image %s x %s cm (recadree pour remplir)" % (to_cm(s.width), to_cm(s.height))
    if s.kind == "chart":
        return "graphique %s x %s cm" % (to_cm(s.width), to_cm(s.height))
    return "tableau %s x %s cm" % (to_cm(s.width), to_cm(s.height))


def summarize(title, body):
    parts = []
    if title:
        parts.append("1 titre")
    for kind, label in (("text", "texte(s)"), ("picture", "image(s)"), ("chart", "graphique(s)"), ("table", "tableau(x)")):
        n = sum(1 for s in body if s.kind == kind)
        if n:
            parts.append("%d %s" % (n, label))
    return ", ".join(parts) or "decor seul"


# ------------------------------------------------------------------ clonage

def clone_slide(prs, src, layout=None):
    """Cree une diapo = copie de src. Retourne (new_slide, mapping {id(el_source): el_copie})."""
    new = prs.slides.add_slide(layout if layout is not None else src.slide_layout)
    tree = new.shapes._spTree
    for child in list(tree):
        if etree.QName(child).localname not in ("nvGrpSpPr", "grpSpPr"):
            tree.remove(child)
    csld_new, csld_src = new._element.find(_QP + "cSld"), src._element.find(_QP + "cSld")
    bg = csld_src.find(_QP + "bg")
    if bg is not None:
        old = csld_new.find(_QP + "bg")
        if old is not None:
            csld_new.remove(old)
        nb = copy.deepcopy(bg)
        remap_rels(nb, src.part, new.part)
        csld_new.insert(0, nb)
    mapping = {}
    for child in src.shapes._spTree:
        if etree.QName(child).localname in ("nvGrpSpPr", "grpSpPr", "extLst"):
            continue
        el = copy.deepcopy(child)
        for a, b in zip(child.iter(), el.iter()):
            mapping[id(a)] = b
        # les graphiques sont reconstruits (evite de partager la partie 'chart' avec le modele)
        gd = el.find(".//" + _QA + "graphicData")
        if etree.QName(el).localname == "graphicFrame" and gd is not None and gd.get("uri") == _CHART_URI:
            for a in child.iter():
                mapping.pop(id(a), None)
            mapping[id(child)] = None
            continue
        remap_rels(el, src.part, new.part)
        renumber_ids(new, el)
        tree.append(el)
    return new, mapping


# ------------------------------------------------------------------ texte

def _a(tag):
    return _QA + tag


def _proto_run_props(p):
    r = p.find(_a("r"))
    if r is not None and r.find(_a("rPr")) is not None:
        return r.find(_a("rPr"))
    e = p.find(_a("endParaRPr"))
    return e


def _new_run(rpr_proto, text, bold, lang):
    r = etree.Element(_a("r"))
    if rpr_proto is not None:
        rpr = copy.deepcopy(rpr_proto)
        rpr.tag = _a("rPr")
    else:
        rpr = etree.Element(_a("rPr"))
    if lang:
        rpr.set("lang", lang)
    if bold:
        rpr.set("b", "1")
    r.append(rpr)
    t = etree.SubElement(r, _a("t"))
    t.text = text
    return r


def _runs(rpr_proto, text, bold_all, lang):
    out = []
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        bold = bold_all
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, bold = part[2:-2], True
        out.append(_new_run(rpr_proto, part, bold, lang))
    return out or [_new_run(rpr_proto, "", bold_all, lang)]


def fill_text_shape(el, items, lang, positional):
    """Remplace le texte d'une forme en reprenant la mise en forme des paragraphes d'exemple.

    items : [(texte, niveau, gras)] ; positional=True : le paragraphe i du texte reprend la mise en
    forme du paragraphe i de l'exemple (ex. intertitre gras puis texte), sinon par niveau.
    """
    tx = el.find(_QP + "txBody")
    if tx is None:
        raise ValueError("forme sans zone de texte")
    protos = [p for p in tx.findall(_a("p")) if "".join(t.text or "" for t in p.iter(_a("t"))).strip()]
    if not protos:
        protos = tx.findall(_a("p"))
    for p in tx.findall(_a("p")):
        tx.remove(p)

    def lvl(p):
        ppr = p.find(_a("pPr"))
        return int(ppr.get("lvl", "0")) if ppr is not None else 0

    for i, (text, level, bold) in enumerate(items):
        if positional:
            proto = protos[min(i, len(protos) - 1)]
        else:
            same = [p for p in protos if lvl(p) == level]
            proto = same[0] if same else min(protos, key=lambda p: abs(lvl(p) - level))
        p = copy.deepcopy(proto)
        rpr = _proto_run_props(proto)
        end = p.find(_a("endParaRPr"))
        for child in list(p):
            if child.tag in (_a("r"), _a("br"), _a("fld")):
                p.remove(child)
        ppr = p.find(_a("pPr"))
        if not positional and level != lvl(proto):
            if ppr is None:
                ppr = etree.Element(_a("pPr"))
                p.insert(0, ppr)
            ppr.set("lvl", str(level))
        runs = _runs(rpr, text, bold, lang)
        for r in runs:
            if end is not None:
                end.addprevious(r)
            else:
                p.append(r)
        tx.append(p)


# ------------------------------------------------------------------ images

def set_picture(pic_el, new_rid, img_w, img_h, box_w, box_h, fit="cover"):
    blip = pic_el.find(".//" + _a("blip"))
    for k in list(blip.attrib):
        if k.startswith(_QR):
            del blip.attrib[k]
    blip.set(_QR + "embed", new_rid)
    sp = pic_el.find(_QP + "blipFill")
    src = sp.find(_a("srcRect"))
    if src is not None:
        sp.remove(src)
    if fit == "cover" and img_w and img_h and box_w and box_h:
        ir, br = img_w / float(img_h), box_w / float(box_h)
        l = r = t = b = 0
        if ir > br:      # image plus large : rogner a gauche/droite
            cut = (1 - br / ir) / 2.0
            l = r = int(cut * 100000)
        elif ir < br:
            cut = (1 - ir / br) / 2.0
            t = b = int(cut * 100000)
        if l or t:
            node = etree.Element(_a("srcRect"))
            for k, v in (("l", l), ("t", t), ("r", r), ("b", b)):
                if v:
                    node.set(k, str(v))
            blip.addnext(node)
