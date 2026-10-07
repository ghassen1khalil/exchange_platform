# -*- coding: utf-8 -*-
"""Bibliotheque d'icones : formes reprises du template + fichiers PNG/JPG."""
from __future__ import annotations

import copy
import difflib
import re
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from .common import ASSETS_DIR, R_NS, ToolError, slugify, to_cm

ICON_DIR = ASSETS_DIR / "icons"
IMAGE_EXT = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff")

_GENERIC = re.compile(r"^(picture|image|group|groupe|graphic|graphique|freeform|forme libre|object|objet|"
                      r"rectangle|oval|ellipse|shape|forme|icon|icone|icône|content placeholder)\s*\d*$", re.I)
_FILELIKE = re.compile(r"\.(png|jpe?g|gif|svg|emf|wmf|bmp|tiff?)$", re.I)
_HINT = re.compile(r"icon|icone|icône|picto|pictogram", re.I)


class IconEntry(object):
    def __init__(self, name, kind, width, height, source, element=None, part=None, path=None):
        self.name = name
        self.kind = kind          # 'shape' | 'file'
        self.width = width        # EMU
        self.height = height
        self.source = source      # texte lisible (ex: "template diapo 12" / "assets/icons/x.png")
        self.element = element
        self.part = part
        self.path = path


def _is_candidate(shape):
    if shape.is_placeholder:
        return False
    st = shape.shape_type
    if st in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.FREEFORM):
        return True
    if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
        return False
    return False


def _alt(shape):
    try:
        c = shape._element.xpath(".//p:cNvPr")[0]
        return (c.get("descr") or "").strip()
    except Exception:
        return ""


def _label_near(shape, shapes):
    """Texte court place juste sous l'icone (legende) - sert a nommer les icones anonymes."""
    best, best_d = None, None
    cx = (shape.left or 0) + (shape.width or 0) / 2.0
    bottom = (shape.top or 0) + (shape.height or 0)
    for s in shapes:
        if s is shape or s.is_placeholder or not s.has_text_frame:
            continue
        txt = s.text_frame.text.strip()
        if not txt or len(txt) > 40 or "\n" in txt:
            continue
        sl, sw = s.left or 0, s.width or 0
        if not (sl - (shape.width or 0) * 0.5 <= cx <= sl + sw + (shape.width or 0) * 0.5):
            continue
        d = (s.top or 0) - bottom
        if d < -(shape.height or 0) * 0.2 or d > (shape.height or 1) * 2.0:
            continue
        if best_d is None or d < best_d:
            best, best_d = txt, d
    return best


class IconLibrary(object):
    def __init__(self):
        self.entries = {}   # name -> IconEntry

    # ---------------------------------------------------------- construction
    @classmethod
    def build(cls, prs, cfg):
        lib = cls()
        lib._scan_template(prs, cfg)
        lib._scan_files()
        return lib

    def _unique(self, base):
        name, i = base, 2
        while name in self.entries:
            name = "%s-%d" % (base, i)
            i += 1
        return name

    def _scan_template(self, prs, cfg):
        mode = cfg.get("icon_library_slides", "auto")
        for n, slide in enumerate(prs.slides, 1):
            if isinstance(mode, list):
                if n not in mode:
                    continue
            elif mode == "all":
                pass
            else:  # auto : diapositives dont le titre ou un texte evoque des icones
                texts = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
                if not _HINT.search(texts):
                    continue
            shapes = list(slide.shapes)
            for sh in shapes:
                if not _is_candidate(sh):
                    continue
                label = _alt(sh) or sh.name or ""
                if not label or _GENERIC.match(label.strip()) or _FILELIKE.search(label.strip()):
                    near = _label_near(sh, shapes)
                    label = near or ("diapo%d-%s" % (n, label or "icone"))
                name = self._unique(slugify(label))
                self.entries[name] = IconEntry(
                    name, "shape", sh.width or 0, sh.height or 0,
                    "template, diapo %d (forme \"%s\")" % (n, sh.name),
                    element=sh._element, part=sh.part)

    def _scan_files(self):
        if not ICON_DIR.is_dir():
            return
        try:
            from PIL import Image
        except Exception:
            Image = None
        for f in sorted(ICON_DIR.iterdir()):
            if f.suffix.lower() not in IMAGE_EXT:
                continue
            w = h = 0
            if Image is not None:
                try:
                    with Image.open(str(f)) as im:
                        w, h = im.size
                except Exception:
                    pass
            name = self._unique(slugify(f.stem))
            # taille native fictive : 96 dpi
            self.entries[name] = IconEntry(name, "file", int(w * 9525), int(h * 9525),
                                           "assets/icons/%s" % f.name, path=f)

    # ---------------------------------------------------------- acces
    def get(self, name):
        key = slugify(name)
        if key in self.entries:
            return self.entries[key]
        close = difflib.get_close_matches(key, list(self.entries), n=4, cutoff=0.4)
        hint = (" Proches : %s." % ", ".join(close)) if close else ""
        avail = ", ".join(sorted(self.entries)) or "(aucune icone disponible)"
        raise ToolError('icone "%s" inconnue.%s Disponibles : %s' % (name, hint, avail))

    def listing(self):
        return sorted(self.entries.values(), key=lambda e: e.name)

    # ---------------------------------------------------------- insertion
    def place(self, slide, name, box, fit="contain"):
        """Insere l'icone dans box=(left, top, width, height) en EMU, centree, ratio conserve."""
        entry = self.get(name)
        left, top, bw, bh = box
        if entry.kind == "file":
            from PIL import Image
            with Image.open(str(entry.path)) as im:
                iw, ih = im.size
            w, h = _fit(iw, ih, bw, bh)
            pic = slide.shapes.add_picture(str(entry.path), Emu(int(left + (bw - w) / 2)),
                                           Emu(int(top + (bh - h) / 2)), Emu(int(w)), Emu(int(h)))
            _set_alt(pic, name)
            return pic
        w, h = _fit(entry.width or 1, entry.height or 1, bw, bh)
        shape = clone_shape(entry, slide)
        shape.left, shape.top = Emu(int(left + (bw - w) / 2)), Emu(int(top + (bh - h) / 2))
        shape.width, shape.height = Emu(int(w)), Emu(int(h))
        _set_alt(shape, name)
        return shape


def _fit(iw, ih, bw, bh):
    if iw <= 0 or ih <= 0:
        return bw, bh
    s = min(bw / float(iw), bh / float(ih))
    return iw * s, ih * s


def _set_alt(shape, text):
    try:
        shape._element.xpath(".//p:cNvPr")[0].set("descr", text)
    except Exception:
        pass


def remap_rels(el, src_part, dst_part):
    """Re-lie les ressources (images, liens...) referencees par `el` vers dst_part."""
    rattr = "{%s}" % R_NS
    for node in el.iter():
        for attr in list(node.attrib):
            if attr.startswith(rattr):
                old = node.get(attr)
                rel = src_part.rels[old]
                if rel.is_external:
                    new = dst_part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
                else:
                    new = dst_part.relate_to(rel.target_part, rel.reltype)
                node.set(attr, new)


def renumber_ids(slide, el):
    next_id = slide.shapes._next_shape_id
    for c in el.xpath(".//p:cNvPr"):
        c.set("id", str(next_id))
        next_id += 1


def clone_shape(entry, slide):
    """Copie profonde d'une forme du template vers `slide`, en re-liant les ressources (images...)."""
    el = copy.deepcopy(entry.element)
    remap_rels(el, entry.part, slide.part)
    renumber_ids(slide, el)
    slide.shapes._spTree.append(el)
    return slide.shapes[-1]


def describe(entry):
    return "%s  (%s x %s cm)  <- %s" % (entry.name, to_cm(entry.width), to_cm(entry.height), entry.source)
