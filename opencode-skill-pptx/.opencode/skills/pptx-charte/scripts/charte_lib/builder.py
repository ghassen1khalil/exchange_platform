# -*- coding: utf-8 -*-
"""Construction d'une presentation a partir d'une spec JSON, en respectant le template."""
from __future__ import annotations

import datetime
import os
import re
from pathlib import Path
from xml.sax.saxutils import escape

from pptx.chart.data import CategoryChartData, XyChartData
from lxml import etree
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import PP_PLACEHOLDER as PH
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Emu, Pt

from .common import (ASSETS_DIR, NSMAP, ToolError, classify_layout, cm, collect_layouts, default_table_style,
                     drop_all_slides, find_template, layout_slots, load_config, open_presentation, read_json,
                     resolve_layout, resolve_path)
from .icons import IconLibrary, _set_alt
from . import pattern as pat
from .textfit import estimate_fill_ratio

SLIDE_KEYS = {"layout", "title", "subtitle", "content", "placeholders", "notes", "footer", "hidden", "extras"}
TOP_KEYS = {"output", "metadata", "slides", "base", "footer", "language"}
BLOCK_KEYS = {"text", "bullets", "image", "icon", "table", "chart", "alt", "fit", "size", "no_bullet", "bold"}
_NSDECL = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
           'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
           'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')

CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "column_stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "column_percent": XL_CHART_TYPE.COLUMN_STACKED_100,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "bar_stacked": XL_CHART_TYPE.BAR_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
    "area_stacked": XL_CHART_TYPE.AREA_STACKED,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
}


class SpecError(Exception):
    pass


class Report(object):
    def __init__(self):
        self.errors = []
        self.warnings = []
        self.lines = []

    def err(self, slide, msg):
        self.errors.append("diapo %s : %s" % (slide, msg))

    def warn(self, slide, msg):
        self.warnings.append("diapo %s : %s" % (slide, msg))


# ------------------------------------------------------------------ texte

_NUM = re.compile(r"^[\s+\-–−]?[\d\s.,]+\s*(%|€|k€|M€|Md€|\$|pts?)?$")


def _flatten_bullets(items, level=0):
    for it in items:
        if isinstance(it, str):
            yield it, level, False
        elif isinstance(it, list):
            for x in _flatten_bullets(it, level + 1):
                yield x
        elif isinstance(it, dict):
            lvl = int(it.get("level", level))
            yield str(it.get("text", "")), lvl, bool(it.get("bold", False))
            if it.get("children"):
                for x in _flatten_bullets(it["children"], lvl + 1):
                    yield x
        else:
            yield str(it), level, False


def _add_runs(p, text, bold_all, lang):
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = bold_all
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, bold = part[2:-2], True
        r = p.add_run()
        r.text = part
        if bold:
            r.font.bold = True
        if lang:
            r._r.get_or_add_rPr().set("lang", lang)


def _fill_text_frame(tf, items, lang, no_bullet=False):
    tf.clear()
    first = True
    for text, level, bold in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if level:
            p.level = min(level, 8)
        if no_bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", "0")
            pPr.set("indent", "0")
            pPr.append(parse_xml("<a:buNone %s/>" % _NSDECL))
        _add_runs(p, text, bold, lang)


def _text_items(value):
    if isinstance(value, str):
        return [(line, 0, False) for line in value.split("\n")]
    raise SpecError("texte attendu (chaine)")


# ------------------------------------------------------------------ zone libre

class FreeSlot(object):
    """Zone calculee pour les layouts sans placeholder de contenu (ex. 'Titre seul')."""

    def __init__(self, left, top, width, height, rank):
        self.idx = -1 - rank
        self.name = "zone libre %d" % (rank + 1)
        self.type_name = "FREE"
        self.kind = "content"
        self.left, self.top, self.width, self.height = left, top, width, height
        self.ph = None


def free_area_slots(prs, by_kind, n_blocks):
    sw, sh = prs.slide_width, prs.slide_height
    titles = by_kind.get("title", [])
    margin_x = cm(1.0)
    left = titles[0].left if titles else margin_x
    width = titles[0].width if titles else sw - 2 * margin_x
    gap = cm(0.4)
    top = (titles[0].top + titles[0].height + gap) if titles else cm(1.5)
    foot_tops = [x.top for k in ("footer", "slidenum", "date") for x in by_kind.get(k, []) if x.height]
    bottom = (min(foot_tops) - gap) if foot_tops else sh - cm(1.2)
    height = max(bottom - top, cm(3))
    n = max(n_blocks, 1)
    col_gap = cm(0.6)
    col_w = int((width - col_gap * (n - 1)) / n)
    return [FreeSlot(left + i * (col_w + col_gap), top, col_w, height, i) for i in range(n)]


# ------------------------------------------------------------------ builder

class DeckBuilder(object):
    def __init__(self, cfg, spec, spec_dir, template_path, base_path=None):
        self.cfg = cfg
        self.spec = spec
        self.spec_dir = spec_dir
        self.report = Report()
        self.lang = spec.get("language") or cfg.get("language") or ""
        self.base_mode = base_path is not None
        self.prs = open_presentation(base_path or template_path)
        self.template_path = template_path
        self.icons = IconLibrary.build(self.prs if not self.base_mode else open_presentation(template_path), cfg)
        tpl_prs = self.prs if not self.base_mode else open_presentation(template_path)
        self.tpl_prs = tpl_prs
        self.template_slides = list(tpl_prs.slides)
        self._pattern_cache = {}
        if not self.base_mode and cfg.get("remove_template_slides", True):
            drop_all_slides(self.prs)
        style = cfg.get("table_style_id", "auto")
        self.table_style = default_table_style(self.prs) if style in (None, "auto") else style
        self.footer_cfg = dict(cfg.get("footer", {}))
        self.footer_cfg.update(spec.get("footer") or {})
        self.limits = cfg["limits"]
        self._footer_warned = False

    # ------------------------------------------------------------ public
    def build(self):
        slides = self.spec.get("slides")
        if not isinstance(slides, list) or not slides:
            raise ToolError("La spec doit contenir une liste 'slides' non vide.")
        for k in self.spec:
            if k not in TOP_KEYS and not k.startswith("_"):
                self.report.warn("-", "cle de premier niveau inconnue ignoree : '%s'" % k)
        for n, s in enumerate(slides, 1):
            try:
                self._build_slide(n, s)
            except SpecError as e:
                self.report.err(n, str(e))
            except ToolError as e:
                self.report.err(n, str(e))
            except Exception as e:  # structure de template inattendue, bug...
                if os.environ.get("PPTX_CHARTE_DEBUG"):
                    raise
                self.report.err(n, "erreur interne (%s: %s). Verifier le layout/les blocs ; "
                                   "PPTX_CHARTE_DEBUG=1 pour le detail" % (type(e).__name__, e))
        self._set_metadata()
        return self.report

    def save(self, out_path):
        out = Path(out_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        try:
            self.prs.save(str(out))
        except PermissionError:
            raise ToolError("Ecriture impossible : %s est probablement ouvert dans PowerPoint. "
                            "Fermez-le ou choisissez un autre nom (ex. -v2)." % out)

    # ------------------------------------------------------------ metadata
    def _set_metadata(self):
        md = self.spec.get("metadata") or {}
        cp = self.prs.core_properties
        cp.title = md.get("title") or cp.title or ""
        cp.author = md.get("author") or self.cfg.get("default_author") or ""
        cp.last_modified_by = cp.author
        cp.subject = md.get("subject") or ""
        cp.keywords = md.get("keywords") or ""
        cp.comments = md.get("comments") or ""
        now = datetime.datetime.now()
        cp.created = now
        cp.modified = now

    # ------------------------------------------------------------ slide
    def _build_slide(self, n, s):
        if not isinstance(s, dict):
            raise SpecError("chaque diapo doit etre un objet JSON")
        unknown = [k for k in s if k not in SLIDE_KEYS and not k.startswith("_")]
        if unknown:
            raise SpecError("cle(s) inconnue(s) %s. Cles permises : %s" % (unknown, sorted(SLIDE_KEYS)))
        ref = s.get("layout")
        if isinstance(ref, str) and ref.strip().lower().startswith("slide:"):
            return self._build_pattern_slide(n, s, ref.strip())
        lref = resolve_layout(self.prs, ref, self.cfg)
        layout = lref.layout
        by_kind, content_slots = layout_slots(layout, self.prs.slide_height)
        role = classify_layout(layout, self.prs.slide_height)
        slide = self.prs.slides.add_slide(layout)
        used = set()
        problems = []

        def ph_of(slot):
            if isinstance(slot, FreeSlot):
                return None
            try:
                return slide.placeholders[slot.idx]
            except KeyError:
                return None

        # titre / sous-titre
        for key, kind in (("title", "title"), ("subtitle", "subtitle")):
            if s.get(key) is None:
                continue
            cands = list(by_kind.get(kind, []))
            if not cands and kind == "subtitle":
                # beaucoup de templates utilisent une zone texte (body) comme sous-titre
                cands = [sl for sl in content_slots if sl.type_name == "BODY" and sl.idx not in used][:1]
            if not cands:
                problems.append("le layout \"%s\" n'a pas de zone '%s' (utilisez un autre layout)" % (layout.name, key))
                continue
            ph = ph_of(cands[0])
            if ph is None or not ph.has_text_frame:
                problems.append("zone '%s' inutilisable" % key)
                continue
            if not isinstance(s[key], str):
                problems.append("'%s' doit etre une chaine" % key)
                continue
            _fill_text_frame(ph.text_frame, [(l, 0, False) for l in s[key].split("\n")], self.lang)
            used.add(cands[0].idx)
            if key == "title" and len(s[key]) > self.limits["max_title_chars"]:
                self.report.warn(n, "titre long (%d car.) : privilegiez un titre plus court" % len(s[key]))

        # placeholders explicites
        explicit = s.get("placeholders") or {}
        all_slots = [sl for lst in by_kind.values() for sl in lst]
        for key, block in explicit.items():
            slot = self._find_slot(all_slots, key)
            if slot is None:
                problems.append("placeholder '%s' introuvable dans \"%s\" (idx/noms : %s)" % (
                    key, layout.name, ", ".join("%s='%s'" % (x.idx, x.name) for x in all_slots)))
                continue
            if slot.idx in used:
                problems.append("placeholder '%s' rempli deux fois" % key)
                continue
            try:
                self._fill_slot(slide, ph_of(slot), slot, block, n)
                used.add(slot.idx)
            except SpecError as e:
                problems.append("placeholder '%s' : %s" % (key, e))

        # contenu sequentiel
        blocks = s.get("content")
        if blocks is not None and not isinstance(blocks, list):
            blocks = [blocks]
        blocks = blocks or []
        free = [sl for sl in content_slots if sl.idx not in used]
        if blocks and not content_slots:
            free = free_area_slots(self.prs, by_kind, len(blocks))
        if len(blocks) > len(free):
            problems.append("%d bloc(s) 'content' mais le layout \"%s\" n'offre que %d zone(s) de contenu libre(s) "
                            "(choisissez un layout adapte)" % (len(blocks), layout.name, len(free)))
            blocks = blocks[:len(free)]
        for slot, block in zip(free, blocks):
            try:
                self._fill_slot(slide, ph_of(slot), slot, block, n)
                used.add(slot.idx)
            except SpecError as e:
                problems.append("content[%d] : %s" % (blocks.index(block), e))

        # zones restees vides -> supprimees
        for slot in content_slots + by_kind.get("subtitle", []) + by_kind.get("title", []):
            if slot.idx in used:
                continue
            ph = ph_of(slot)
            if ph is None:
                continue
            ph._element.getparent().remove(ph._element)
            if slot.kind == "title":
                self.report.warn(n, "aucun titre fourni : zone titre supprimee (accessibilite/navigation)")
            else:
                self.report.warn(n, "zone '%s' (idx %s) du layout \"%s\" laissee vide : supprimee" % (
                    slot.name, slot.idx, layout.name))

        # extras libres
        for ex in (s.get("extras") or []):
            try:
                self._add_extra(slide, ex)
            except (SpecError, ToolError) as e:
                problems.append("extras : %s" % e)

        # pied de page, numero, date
        self._add_footers(slide, layout, role, s.get("footer"), n)

        # notes, masquage
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])
        if s.get("hidden"):
            slide._element.set("show", "0")

        # controle de debordement
        for shp in slide.shapes:
            if shp.has_text_frame and shp.text_frame.text.strip():
                r = estimate_fill_ratio(shp, self.limits["avg_char_width_em"], self.limits["line_height_em"])
                if r > 1.08:
                    self.report.warn(n, "debordement probable dans '%s' (~%d%% de la zone) : raccourcir ou scinder" % (
                        shp.name, int(r * 100)))
        if problems:
            raise SpecError(problems[0] if len(problems) == 1 else "\n    - " + "\n    - ".join(problems))
        title = ""
        for shp in slide.shapes:
            if shp.is_placeholder and shp.placeholder_format.type in (PH.TITLE, PH.CENTER_TITLE) and shp.has_text_frame:
                title = shp.text_frame.text.replace("\n", " ")
        self.report.lines.append("  %2d. [%s] %s" % (n, layout.name, title))

    # ------------------------------------------------------------ diapos modeles
    def _pattern_layout(self, src):
        if not self.base_mode:
            return src.slide_layout
        name = src.slide_layout.name
        for lr in collect_layouts(self.prs):
            if lr.name == name:
                return lr.layout
        raise SpecError("layout \"%s\" de la diapo modele introuvable dans la presentation de base" % name)

    def _pattern_slots(self, num):
        if num not in self._pattern_cache:
            src = self.template_slides[num - 1]
            self._pattern_cache[num] = pat.analyze_slide(src, self.tpl_prs.slide_width, self.tpl_prs.slide_height)
        return self._pattern_cache[num]

    def _build_pattern_slide(self, n, s, ref):
        try:
            num = int(ref.split(":", 1)[1])
        except ValueError:
            raise SpecError("reference de diapo modele invalide : '%s' (attendu : slide:N)" % ref)
        if num < 1 or num > len(self.template_slides):
            raise SpecError("diapo modele %d inexistante (le template en contient %d)" % (num, len(self.template_slides)))
        title_slot, body = self._pattern_slots(num)
        src = self.template_slides[num - 1]
        new, mapping = pat.clone_slide(self.prs, src, self._pattern_layout(src))
        problems, used = [], set()

        def fill(slot, block, label):
            try:
                self._fill_pslot(new, mapping, slot, block, n)
                used.add(id(slot))
            except SpecError as e:
                problems.append("%s : %s" % (label, e))

        if s.get("title") is not None:
            if title_slot is None:
                problems.append("le modele slide:%d n'a pas de zone titre" % num)
            else:
                fill(title_slot, {"text": str(s["title"])}, "title")
        if s.get("subtitle") is not None:
            problems.append("'subtitle' n'existe pas en mode diapo modele : utilisez 'content' (voir le catalogue)")

        explicit = s.get("placeholders") or {}
        for key, block in explicit.items():
            hit = [sl for sl in ([title_slot] if title_slot else []) + body if sl.name.lower() == str(key).lower()]
            if not hit:
                problems.append("zone '%s' introuvable dans slide:%d (noms : %s)" % (
                    key, num, ", ".join(sl.name for sl in body)))
                continue
            fill(hit[0], block, "placeholders['%s']" % key)

        blocks = s.get("content")
        if blocks is not None and not isinstance(blocks, list):
            blocks = [blocks]
        blocks = blocks or []
        free = [sl for sl in body if id(sl) not in used]
        if len(blocks) > len(free):
            problems.append("%d bloc(s) 'content' mais slide:%d n'offre que %d zone(s) libre(s)" % (len(blocks), num, len(free)))
            blocks = blocks[:len(free)]
        for i, (slot, block) in enumerate(zip(free, blocks)):
            fill(slot, block, "content[%d]" % i)

        # zones non remplies : supprimees (jamais de lorem ipsum residuel)
        for slot in ([title_slot] if title_slot else []) + body:
            if id(slot) in used:
                continue
            el = mapping.get(id(slot.el))
            if el is not None and el.getparent() is not None:
                el.getparent().remove(el)
            if slot.kind != "title":
                self.report.warn(n, "zone '%s' (%s) de slide:%d laissee vide : supprimee (ses elements decoratifs - numero, picto - "
                                    "restent : preferer un modele avec le bon nombre de zones)" % (slot.name, slot.kind, num))
            else:
                self.report.warn(n, "aucun titre fourni : zone titre supprimee")

        for ex in (s.get("extras") or []):
            try:
                self._add_extra(new, ex)
            except (SpecError, ToolError) as e:
                problems.append("extras : %s" % e)
        if s.get("notes"):
            new.notes_slide.notes_text_frame.text = str(s["notes"])
        if s.get("hidden"):
            new._element.set("show", "0")
        for shp in new.shapes:
            if shp.has_text_frame and shp.text_frame.text.strip():
                r = estimate_fill_ratio(shp, self.limits["avg_char_width_em"], self.limits["line_height_em"])
                if r > 1.08:
                    self.report.warn(n, "debordement probable dans '%s' (~%d%% de la zone) : raccourcir ou scinder" % (
                        shp.name, int(r * 100)))
        if problems:
            raise SpecError(problems[0] if len(problems) == 1 else "\n    - " + "\n    - ".join(problems))
        title = ""
        if title_slot is not None and s.get("title"):
            title = str(s["title"]).replace("\n", " ")
        self.report.lines.append("  %2d. [slide:%d] %s" % (n, num, title))

    def _fill_pslot(self, new, mapping, slot, block, n):
        el = mapping.get(id(slot.el))
        if isinstance(block, str):
            block = {"text": block}
        elif isinstance(block, list):
            block = {"bullets": block}
        if not isinstance(block, dict):
            raise SpecError("bloc invalide (chaine, liste ou objet attendu)")
        unk = [k for k in block if k not in BLOCK_KEYS and not k.startswith("_")]
        if unk:
            raise SpecError("cle(s) inconnue(s) %s dans un bloc. Permises : %s" % (unk, sorted(BLOCK_KEYS)))
        kinds = [k for k in ("text", "bullets", "image", "icon", "table", "chart") if k in block]
        if len(kinds) != 1:
            raise SpecError("un bloc doit contenir exactement un type parmi text/bullets/image/icon/table/chart (recu : %s)" % kinds)
        kind = kinds[0]
        box = (slot.left, slot.top, slot.width, slot.height)

        if kind in ("text", "bullets"):
            if slot.kind not in ("text", "title"):
                raise SpecError("la zone '%s' est de type %s : elle n'accepte pas de texte" % (slot.name, slot.kind))
            if kind == "text":
                items = _text_items(block["text"])
                if block.get("bold"):
                    items = [(t, l, True) for t, l, _ in items]
            else:
                if not isinstance(block["bullets"], list):
                    raise SpecError("'bullets' doit etre une liste")
                items = list(_flatten_bullets(block["bullets"]))
                nb = sum(1 for _, lvl, _ in items if lvl == 0)
                if nb > self.limits["max_bullets_per_placeholder"]:
                    self.report.warn(n, "%d puces de niveau 1 (max conseille %d) : scinder la diapo" % (
                        nb, self.limits["max_bullets_per_placeholder"]))
            pat.fill_text_shape(el, items, self.lang, positional=(kind == "text"))
            return

        if slot.kind in ("title",):
            raise SpecError("la zone titre n'accepte que du texte")
        anchor = el
        if kind == "image":
            path = resolve_path(block["image"], self.spec_dir, ASSETS_DIR)
            if path is None:
                raise SpecError("image introuvable : %s" % block["image"])
            if not block.get("alt"):
                self.report.warn(n, "image '%s' sans texte alternatif ('alt')" % Path(str(block["image"])).name)
            from PIL import Image
            with Image.open(str(path)) as im:
                iw, ih = im.size
            fit = block.get("fit", "cover")
            if el is not None and etree.QName(el).localname == "pic":
                _, rid = new.part.get_or_add_image_part(str(path))
                pat.set_picture(el, rid, iw, ih, slot.width, slot.height, fit)
                _set_alt(el, block.get("alt", ""))
                return
            if fit == "contain":
                pic = self._add_contained_picture(new, path, box)
            else:
                pic = new.shapes.add_picture(str(path), Emu(box[0]), Emu(box[1]), Emu(box[2]), Emu(box[3]))
                _, rid = new.part.get_or_add_image_part(str(path))
                pat.set_picture(pic._element, rid, iw, ih, box[2], box[3], "cover")
            _set_alt(pic, block.get("alt", ""))
            new_el = pic._element
        elif kind == "icon":
            size = block.get("size")
            b = box
            if size:
                side = cm(size)
                b = (box[0] + (box[2] - side) // 2, box[1] + (box[3] - side) // 2, side, side)
            new_el = self.icons.place(new, block["icon"], b)._element
        elif kind == "table":
            style = None
            if slot.kind == "table" and el is not None:
                sid = el.xpath(".//a:tableStyleId")
                style = sid[0].text if sid else None
            new_el = self._add_table(new, block["table"], box, n, style_id=style)._element
        else:
            new_el = self._add_chart(new, block["chart"], box, block.get("alt", ""), n)._element
        if anchor is not None and anchor.getparent() is not None:
            anchor.addprevious(new_el)
            anchor.getparent().remove(anchor)

    # ------------------------------------------------------------ slots
    @staticmethod
    def _find_slot(slots, key):
        k = str(key).strip()
        for sl in slots:
            if str(sl.idx) == k.replace("idx:", "").strip():
                return sl
        for sl in slots:
            if sl.name.lower() == k.lower():
                return sl
        return None

    def _fill_slot(self, slide, ph, slot, block, n):
        if ph is None and not isinstance(slot, FreeSlot):
            raise SpecError("placeholder absent de la diapo")
        if isinstance(block, str):
            block = {"text": block}
        elif isinstance(block, list):
            block = {"bullets": block}
        if not isinstance(block, dict):
            raise SpecError("bloc invalide (chaine, liste ou objet attendu)")
        unk = [k for k in block if k not in BLOCK_KEYS and not k.startswith("_")]
        if unk:
            raise SpecError("cle(s) inconnue(s) %s dans un bloc. Permises : %s" % (unk, sorted(BLOCK_KEYS)))
        kinds = [k for k in ("text", "bullets", "image", "icon", "table", "chart") if k in block]
        if len(kinds) != 1:
            raise SpecError("un bloc doit contenir exactement un type parmi text/bullets/image/icon/table/chart (recu : %s)" % kinds)
        kind = kinds[0]
        if ph is not None:
            box = (ph.left or 0, ph.top or 0, ph.width or 0, ph.height or 0)
        else:
            box = (slot.left, slot.top, slot.width, slot.height)

        if kind in ("text", "bullets"):
            if ph is None:
                raise SpecError("texte impossible dans une zone libre : choisissez un layout avec une zone de texte")
            if slot.type_name in ("PICTURE", "BITMAP", "TABLE", "CHART", "MEDIA_CLIP", "ORG_CHART"):
                raise SpecError("la zone '%s' est de type %s : elle n'accepte pas de texte (utilisez image/table/chart)" % (
                    slot.name, slot.type_name))
            if not ph.has_text_frame:
                raise SpecError("la zone '%s' (type %s) n'accepte pas de texte" % (slot.name, slot.type_name))
            if kind == "text":
                items = _text_items(block["text"])
                if block.get("bold"):
                    items = [(t, l, True) for t, l, _ in items]
            else:
                if not isinstance(block["bullets"], list):
                    raise SpecError("'bullets' doit etre une liste")
                items = list(_flatten_bullets(block["bullets"]))
                nb = sum(1 for _, lvl, _ in items if lvl == 0)
                if nb > self.limits["max_bullets_per_placeholder"]:
                    self.report.warn(n, "%d puces de niveau 1 (max conseille %d) : scinder la diapo" % (
                        nb, self.limits["max_bullets_per_placeholder"]))
            _fill_text_frame(ph.text_frame, items, self.lang, no_bullet=bool(block.get("no_bullet")))
            return

        if kind == "image":
            path = resolve_path(block["image"], self.spec_dir, ASSETS_DIR)
            if path is None:
                raise SpecError("image introuvable : %s" % block["image"])
            if not block.get("alt"):
                self.report.warn(n, "image '%s' sans texte alternatif ('alt')" % Path(str(block["image"])).name)
            fit = block.get("fit")
            if ph is not None and hasattr(ph, "insert_picture") and fit != "contain":
                pic = ph.insert_picture(str(path))
                _set_alt(pic, block.get("alt", ""))
                return
            pic = self._add_contained_picture(slide, path, box)
            _set_alt(pic, block.get("alt", ""))
            self._remove(ph)
            return

        if kind == "icon":
            size = block.get("size")
            b = box
            if size:
                side = cm(size)
                b = (box[0] + (box[2] - side) // 2, box[1] + (box[3] - side) // 2, side, side)
            self.icons.place(slide, block["icon"], b)
            self._remove(ph)
            return

        if kind == "table":
            self._add_table(slide, block["table"], box, n)
            self._remove(ph)
            return

        if kind == "chart":
            self._add_chart(slide, block["chart"], box, block.get("alt", ""), n)
            self._remove(ph)
            return

    @staticmethod
    def _remove(ph):
        if ph is not None:
            ph._element.getparent().remove(ph._element)

    # ------------------------------------------------------------ images
    @staticmethod
    def _add_contained_picture(slide, path, box):
        from PIL import Image
        with Image.open(str(path)) as im:
            iw, ih = im.size
        left, top, bw, bh = box
        s = min(bw / float(iw), bh / float(ih))
        w, h = iw * s, ih * s
        return slide.shapes.add_picture(str(path), Emu(int(left + (bw - w) / 2)), Emu(int(top + (bh - h) / 2)),
                                        Emu(int(w)), Emu(int(h)))

    # ------------------------------------------------------------ table
    def _add_table(self, slide, t, box, n, style_id=None):
        if not isinstance(t, dict) or "rows" not in t:
            raise SpecError("table : objet avec 'rows' (et 'header' optionnel) attendu")
        header = t.get("header")
        rows = t["rows"]
        if not rows or not isinstance(rows, list):
            raise SpecError("table : 'rows' vide")
        ncols = len(header) if header else len(rows[0])
        for r in rows:
            if len(r) != ncols:
                raise SpecError("table : toutes les lignes doivent avoir %d colonnes" % ncols)
        nrows = len(rows) + (1 if header else 0)
        font_pt = float(t.get("font_pt", self.cfg.get("table_font_pt", 12)))
        left, top, bw, bh = box
        row_h = min(int(bh / nrows), int(Pt(font_pt) * 2.6))
        gf = slide.shapes.add_table(nrows, ncols, Emu(left), Emu(top), Emu(bw), Emu(row_h * nrows))
        tbl = gf.table
        gf._element.xpath(".//a:tableStyleId")[0].text = style_id or self.table_style
        tbl.first_row = bool(header)
        widths = t.get("col_widths") or [1] * ncols
        tot = float(sum(widths))
        for i, wv in enumerate(widths[:ncols]):
            tbl.columns[i].width = Emu(int(bw * wv / tot))
        for r in tbl.rows:
            r.height = Emu(row_h)
        aligns = t.get("align")
        if not aligns:
            aligns = []
            for c in range(ncols):
                col = [str(r[c]) for r in rows]
                aligns.append("right" if all(_NUM.match(x or "") for x in col if x) and any(col) else "left")
        amap = {"l": PP_ALIGN.LEFT, "left": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "right": PP_ALIGN.RIGHT,
                "c": PP_ALIGN.CENTER, "center": PP_ALIGN.CENTER}
        data = ([header] if header else []) + rows
        for ri, row in enumerate(data):
            for ci, val in enumerate(row):
                cell = tbl.cell(ri, ci)
                tf = cell.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                _add_runs(p, "" if val is None else str(val), False, self.lang)
                for run in p.runs:
                    run.font.size = Pt(font_pt)
                al = amap.get(str(aligns[ci]).lower())
                if al is not None:
                    p.alignment = al
        alt = t.get("alt", "")
        gf._element.xpath(".//p:cNvPr")[0].set("descr", alt or "Tableau")
        return gf

    # ------------------------------------------------------------ chart
    def _add_chart(self, slide, c, box, alt, n):
        if not isinstance(c, dict):
            raise SpecError("chart : objet attendu")
        ctype = c.get("type", "column")
        if ctype not in CHART_TYPES:
            raise SpecError("chart.type '%s' inconnu (permis : %s)" % (ctype, ", ".join(sorted(CHART_TYPES))))
        series = c.get("series")
        if not series:
            raise SpecError("chart.series manquant")
        if ctype == "scatter":
            data = XyChartData()
            for s in series:
                ser = data.add_series(s.get("name", ""))
                for x, y in s["points"]:
                    ser.add_data_point(x, y)
        else:
            cats = c.get("categories")
            if not cats:
                raise SpecError("chart.categories manquant")
            data = CategoryChartData()
            data.categories = cats
            for s in series:
                if len(s["values"]) != len(cats):
                    raise SpecError("chart : la serie '%s' a %d valeurs pour %d categories" % (
                        s.get("name", "?"), len(s["values"]), len(cats)))
                data.add_series(s.get("name", ""), s["values"])
        left, top, bw, bh = box
        gf = slide.shapes.add_chart(CHART_TYPES[ctype], Emu(left), Emu(top), Emu(bw), Emu(bh), data)
        chart = gf.chart
        chart.font.size = Pt(float(c.get("font_pt", self.cfg.get("chart_font_pt", 12))))
        if c.get("title"):
            chart.has_title = True
            chart.chart_title.text_frame.text = c["title"]
        else:
            chart.has_title = False
        legend = c.get("legend")
        if legend is None:
            legend = len(series) > 1 or ctype in ("pie", "doughnut")
        chart.has_legend = bool(legend)
        if legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        if c.get("data_labels"):
            plot = chart.plots[0]
            plot.has_data_labels = True
            if c.get("number_format"):
                plot.data_labels.number_format = c["number_format"]
                plot.data_labels.number_format_is_linked = False
        if ctype in ("pie", "doughnut"):
            chart.plots[0].vary_by_categories = True
        if ctype == "scatter":
            from pptx.enum.chart import XL_MARKER_STYLE
            for ser in chart.plots[0].series:
                ser.marker.style = XL_MARKER_STYLE.CIRCLE
                ser.marker.size = 9
                ser.format.line.fill.background()
        gf._element.xpath(".//p:cNvPr")[0].set("descr", alt or c.get("title") or "Graphique")
        return gf

    # ------------------------------------------------------------ extras
    def _add_extra(self, slide, ex):
        if not isinstance(ex, dict):
            raise SpecError("extra invalide")
        for req in ("x", "y"):
            if req not in ex:
                raise SpecError("extra : '%s' (cm) requis" % req)
        x, y = cm(ex["x"]), cm(ex["y"])
        if "icon" in ex:
            size = cm(ex.get("size", 1.5))
            w, h = (cm(ex["w"]), cm(ex["h"])) if "w" in ex and "h" in ex else (size, size)
            self.icons.place(slide, ex["icon"], (x, y, w, h))
        elif "image" in ex:
            path = resolve_path(ex["image"], self.spec_dir, ASSETS_DIR)
            if path is None:
                raise SpecError("image introuvable : %s" % ex["image"])
            if "w" not in ex or "h" not in ex:
                raise SpecError("extra image : 'w' et 'h' (cm) requis")
            pic = self._add_contained_picture(slide, path, (x, y, cm(ex["w"]), cm(ex["h"])))
            _set_alt(pic, ex.get("alt", ""))
        elif "text" in ex:
            if "w" not in ex or "h" not in ex:
                raise SpecError("extra texte : 'w' et 'h' (cm) requis")
            tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(cm(ex["w"])), Emu(cm(ex["h"])))
            tb.text_frame.word_wrap = True
            _fill_text_frame(tb.text_frame, _text_items(str(ex["text"])), self.lang)
            size = ex.get("size_pt", self.cfg.get("extra_text_pt", 18))
            for p in tb.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(float(size))
                    if ex.get("bold"):
                        r.font.bold = True
                    if ex.get("color"):
                        self._theme_color(r.font.color, ex["color"])
                if ex.get("align"):
                    p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT,
                                   "center": PP_ALIGN.CENTER}.get(ex["align"], PP_ALIGN.LEFT)
        else:
            raise SpecError("extra : 'icon', 'image' ou 'text' requis")

    @staticmethod
    def _theme_color(color_fmt, name):
        from pptx.enum.dml import MSO_THEME_COLOR as TC
        mp = {"dk1": TC.DARK_1, "lt1": TC.LIGHT_1, "dk2": TC.DARK_2, "lt2": TC.LIGHT_2,
              "text1": TC.TEXT_1, "background1": TC.BACKGROUND_1, "text2": TC.TEXT_2, "background2": TC.BACKGROUND_2,
              "accent1": TC.ACCENT_1, "accent2": TC.ACCENT_2, "accent3": TC.ACCENT_3,
              "accent4": TC.ACCENT_4, "accent5": TC.ACCENT_5, "accent6": TC.ACCENT_6}
        if name not in mp:
            raise SpecError("couleur '%s' inconnue : seules les couleurs du theme sont permises (%s)" % (
                name, ", ".join(sorted(mp))))
        color_fmt.theme_color = mp[name]

    # ------------------------------------------------------------ footers
    def _add_footers(self, slide, layout, role, slide_footer, n):
        cfgf = self.footer_cfg
        if slide_footer is False or role in cfgf.get("hide_on_roles", []):
            return
        text = slide_footer if isinstance(slide_footer, str) else cfgf.get("text", "")
        want = {"sldNum": bool(cfgf.get("slide_number", True)), "ftr": bool(text), "dt": bool(cfgf.get("date", False))}
        lang = self.lang or "fr-FR"
        next_id = slide.shapes._next_shape_id
        for el in slide.shapes._spTree.iterfind(".//p:nvPr/p:ph", NSMAP):
            if el.get("type") in want:
                want[el.get("type")] = False
        for ph in layout.placeholders:
            lph = ph._element.find(".//p:nvPr/p:ph", NSMAP)
            t = lph.get("type") if lph is not None else None
            if t not in want or not want[t]:
                continue
            idx = lph.get("idx", "0")
            sz = lph.get("sz")
            szattr = ' sz="%s"' % sz if sz else ""
            if t == "sldNum":
                body = ('<a:p><a:fld id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum">'
                        '<a:rPr lang="%s"/><a:t>‹#›</a:t></a:fld><a:endParaRPr lang="%s"/></a:p>' % (lang, lang))
                nm = "Numero de diapositive"
            elif t == "ftr":
                body = '<a:p><a:r><a:rPr lang="%s"/><a:t>%s</a:t></a:r></a:p>' % (lang, escape(text))
                nm = "Pied de page"
            else:
                body = '<a:p><a:r><a:rPr lang="%s"/><a:t>%s</a:t></a:r></a:p>' % (
                    lang, datetime.date.today().strftime("%d/%m/%Y"))
                nm = "Date"
            xml = ('<p:sp %s><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
                   '<p:nvPr><p:ph type="%s"%s idx="%s"/></p:nvPr></p:nvSpPr><p:spPr/>'
                   '<p:txBody><a:bodyPr/><a:lstStyle/>%s</p:txBody></p:sp>' % (_NSDECL, next_id, nm, t, szattr, idx, body))
            slide.shapes._spTree.append(parse_xml(xml))
            next_id += 1
            want[t] = False
        if want["sldNum"] and not self._footer_warned:
            self._footer_warned = True
            self.report.warn(n, "le layout \"%s\" n'a pas de zone numero de diapo : numerotation absente" % layout.name)


# ------------------------------------------------------------------ API

def build_from_spec(spec_path, output=None, template=None, base=None):
    cfg = load_config()
    spec_path = Path(spec_path)
    spec = read_json(spec_path)
    if not isinstance(spec, dict):
        raise ToolError("La spec doit etre un objet JSON.")
    tpl = find_template(template, cfg)
    base_p = Path(base) if base else None
    if base_p is not None and not base_p.is_file():
        raise ToolError("Presentation de base introuvable : %s" % base)
    b = DeckBuilder(cfg, spec, spec_path.resolve().parent, tpl, str(base_p) if base_p else None)
    report = b.build()
    out = output or spec.get("output") or str(Path("output") / (spec_path.stem.replace(".spec", "") + ".pptx"))
    return b, report, Path(out), tpl
