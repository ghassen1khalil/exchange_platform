# -*- coding: utf-8 -*-
"""Utilitaires communs : configuration, template, theme, layouts."""
from __future__ import annotations

import difflib
import io
import json
import os
import re
import sys
import unicodedata
import zipfile
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parents[2]
ASSETS_DIR = SKILL_DIR / "assets"
CONFIG_PATH = SKILL_DIR / "config" / "charte.json"

CT_POTX = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
CT_PPSX = "application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml"
CT_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NSMAP = {"a": A_NS, "p": P_NS, "r": R_NS}

EMU_PER_CM = 360000
EMU_PER_PT = 12700

THEME_COLOR_ORDER = ["dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3",
                     "accent4", "accent5", "accent6", "hlink", "folHlink"]

DEFAULT_CONFIG = {
    "template": None,
    "language": "fr-FR",
    "default_author": "",
    "remove_template_slides": True,
    "layout_aliases": {},
    "allowed_layouts": None,
    "forbidden_layouts": [],
    "icon_library_slides": "auto",
    "extra_palette": [],
    "table_style_id": "auto",
    "table_font_pt": 14,
    "chart_font_pt": 14,
    "min_font_pt": 12,
    "extra_text_pt": 18,
    "footer": {
        "text": "",
        "slide_number": True,
        "date": False,
        "hide_on_roles": ["cover", "section", "end"],
    },
    "limits": {
        "max_bullets_per_placeholder": 7,
        "max_words_per_slide": 110,
        "max_title_chars": 90,
        "avg_char_width_em": 0.52,
        "line_height_em": 1.2,
    },
    "forbidden_texts": ["lorem ipsum", "cliquez pour", "click to edit", "cliquez ici"],
}


class ToolError(Exception):
    """Erreur attendue, affichee sans traceback."""


def ensure_utf8_console():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


def slugify(text: str) -> str:
    t = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode("ascii")
    t = re.sub(r"[^a-zA-Z0-9]+", "-", t).strip("-").lower()
    return t or "x"


def read_json(path) -> dict:
    p = Path(path)
    try:
        return json.loads(p.read_text(encoding="utf-8-sig"))
    except FileNotFoundError:
        raise ToolError("Fichier introuvable : %s" % p)
    except json.JSONDecodeError as e:
        raise ToolError("JSON invalide dans %s (ligne %d, colonne %d) : %s" % (p, e.lineno, e.colno, e.msg))


def _deep_merge(base: dict, over: dict) -> dict:
    out = dict(base)
    for k, v in over.items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = _deep_merge(out[k], v)
        else:
            out[k] = v
    return out


def load_config() -> dict:
    cfg = DEFAULT_CONFIG
    if CONFIG_PATH.exists():
        user = read_json(CONFIG_PATH)
        user = {k: v for k, v in user.items() if not k.startswith("_")}
        cfg = _deep_merge(DEFAULT_CONFIG, user)
    return cfg


def find_template(cli_path=None, cfg=None) -> Path:
    cfg = cfg or load_config()
    candidates = []
    if cli_path:
        candidates.append(Path(cli_path).expanduser())
    env = os.environ.get("PPTX_CHARTE_TEMPLATE")
    if env:
        candidates.append(Path(env).expanduser())
    if cfg.get("template"):
        t = Path(cfg["template"]).expanduser()
        candidates.append(t if t.is_absolute() else SKILL_DIR / t)
    for name in ("template.potx", "template.pptx", "modele.potx", "modele.pptx"):
        candidates.append(ASSETS_DIR / name)
    for c in candidates:
        if c.is_file():
            return c.resolve()
    # dernier recours : un unique .potx/.pptx dans assets/
    found = sorted(list(ASSETS_DIR.glob("*.potx")) + list(ASSETS_DIR.glob("*.pptx")))
    if len(found) == 1:
        return found[0].resolve()
    if cli_path or env:
        raise ToolError("Template introuvable : %s" % (cli_path or env))
    raise ToolError(
        "Aucun template trouve. Copiez le modele PowerPoint officiel de l'organisation dans\n"
        "  %s\nsous le nom 'template.potx' (ou template.pptx), puis lancez : run catalog" % ASSETS_DIR)


def _normalize_template_bytes(data: bytes) -> bytes:
    """python-pptx refuse .potx/.ppsx : on corrige le content-type en memoire."""
    try:
        zin = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ToolError("Le fichier n'est pas un document Office Open XML valide (zip corrompu ?).")
    ct = zin.read("[Content_Types].xml").decode("utf-8")
    if "macroEnabled" in ct:
        raise ToolError("Les modeles avec macros (.potm/.pptm) ne sont pas supportes : "
                        "enregistrez le modele en .potx ou .pptx (sans macros).")
    if CT_POTX not in ct and CT_PPSX not in ct:
        return data
    ct = ct.replace(CT_POTX, CT_PPTX).replace(CT_PPSX, CT_PPTX)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                payload = ct.encode("utf-8")
            zout.writestr(item, payload)
    return out.getvalue()


def open_presentation(path):
    from pptx import Presentation
    p = Path(path)
    if not p.is_file():
        raise ToolError("Fichier introuvable : %s" % p)
    data = _normalize_template_bytes(p.read_bytes())
    try:
        return Presentation(io.BytesIO(data))
    except Exception as e:  # pragma: no cover
        raise ToolError("Impossible d'ouvrir %s : %s" % (p.name, e))


# ---------------------------------------------------------------- theme

def theme_info(prs, master_index=0) -> dict:
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    master = prs.slide_masters[master_index]
    part = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    colors = {}
    scheme = root.find(".//a:clrScheme", NSMAP)
    for name in THEME_COLOR_ORDER:
        el = scheme.find("a:%s" % name, NSMAP) if scheme is not None else None
        if el is None or len(el) == 0:
            continue
        child = el[0]
        tag = etree.QName(child).localname
        hexv = child.get("val") if tag == "srgbClr" else child.get("lastClr")
        if hexv:
            colors[name] = hexv.upper()
    major = root.find(".//a:fontScheme/a:majorFont/a:latin", NSMAP)
    minor = root.find(".//a:fontScheme/a:minorFont/a:latin", NSMAP)
    return {
        "color_scheme_name": scheme.get("name") if scheme is not None else "",
        "colors": colors,
        "major_font": major.get("typeface") if major is not None else "",
        "minor_font": minor.get("typeface") if minor is not None else "",
        "theme_name": root.get("name", ""),
    }


def default_table_style(prs) -> str:
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    try:
        part = prs.part.part_related_by(RT.TABLE_STYLES)
        root = etree.fromstring(part.blob)
        return root.get("def") or "{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"
    except Exception:
        return "{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"


# ---------------------------------------------------------------- layouts

class LayoutRef(object):
    def __init__(self, index, master_index, layout):
        self.index = index
        self.master_index = master_index
        self.layout = layout

    @property
    def name(self):
        return self.layout.name


def collect_layouts(prs):
    out = []
    n = 0
    for mi, master in enumerate(prs.slide_masters):
        for layout in master.slide_layouts:
            n += 1
            out.append(LayoutRef(n, mi, layout))
    return out


def ph_kind(ph_type) -> str:
    from pptx.enum.shapes import PP_PLACEHOLDER as T
    if ph_type in (T.TITLE, T.CENTER_TITLE, T.VERTICAL_TITLE):
        return "title"
    if ph_type == T.SUBTITLE:
        return "subtitle"
    if ph_type == T.DATE:
        return "date"
    if ph_type == T.FOOTER:
        return "footer"
    if ph_type == T.SLIDE_NUMBER:
        return "slidenum"
    if ph_type == T.HEADER:
        return "header"
    return "content"


class Slot(object):
    """Un placeholder d'un layout, vu comme une zone remplissable."""

    def __init__(self, ph):
        self.ph = ph
        self.idx = ph.placeholder_format.idx
        self.type = ph.placeholder_format.type
        self.type_name = str(self.type).split(".")[-1].split(" ")[0] if self.type is not None else "OBJECT"
        self.name = ph.name
        self.kind = ph_kind(self.type)
        self.left = ph.left or 0
        self.top = ph.top or 0
        self.width = ph.width or 0
        self.height = ph.height or 0
        try:
            self.prompt = ph.text_frame.text.strip() if ph.has_text_frame else ""
        except Exception:
            self.prompt = ""


def layout_slots(layout, slide_height=None):
    """Retourne (slots_par_kind, content_slots_ordonnes)."""
    slots = [Slot(ph) for ph in layout.placeholders]
    by_kind = {}
    for s in slots:
        by_kind.setdefault(s.kind, []).append(s)
    content = list(by_kind.get("content", []))
    content = order_reading(content, slide_height)
    return by_kind, content


def order_reading(slots, slide_height=None):
    if not slots:
        return []
    tol = (slide_height or 6858000) * 0.04
    ordered = sorted(slots, key=lambda s: (s.top, s.left))
    rows, cur, row_top = [], [], None
    for s in ordered:
        if row_top is None or s.top - row_top <= tol:
            cur.append(s)
            row_top = s.top if row_top is None else row_top
        else:
            rows.append(cur)
            cur, row_top = [s], s.top
    if cur:
        rows.append(cur)
    out = []
    for r in rows:
        out.extend(sorted(r, key=lambda s: s.left))
    return out


_ROLE_TOKENS = [
    ("end", {"merci", "thanks", "thank", "closing", "fin", "end", "contact"}),
    ("cover", {"cover", "couverture", "garde"}),
    ("section", {"section", "chapitre", "chapter", "divider", "intercalaire", "separateur", "transition"}),
    ("agenda", {"agenda", "sommaire", "plan"}),
    ("comparison", {"comparison", "comparaison"}),
]


def classify_layout(layout, slide_height=None) -> str:
    by_kind, content = layout_slots(layout, slide_height)
    slug = slugify(layout.name)
    tokens = set(slug.split("-"))
    for role, kws in _ROLE_TOKENS:
        if tokens & kws:
            return role
    from pptx.enum.shapes import PP_PLACEHOLDER as T
    titles = by_kind.get("title", [])
    if any(s.type == T.CENTER_TITLE for s in titles) or slug.startswith("title-slide") \
            or slug in ("titre", "title", "diapositive-de-titre", "diapo-titre", "page-de-titre"):
        return "cover"
    if not titles and not content:
        return "blank"
    if tokens & {"blank", "vide", "blanc"}:
        return "blank"
    if titles and not content:
        return "title_only"
    n = len(content)
    if n == 1:
        return "picture" if content[0].type in (T.PICTURE, T.BITMAP) else "content"
    if n == 2:
        return "two_content"
    if n >= 3:
        return "three_content"
    return "other"


def _fold(s):
    return slugify(s)


def resolve_layout(prs, ref, cfg=None):
    """Resout un nom / un @role / un #index en LayoutRef (ou leve ToolError)."""
    cfg = cfg or load_config()
    layouts = collect_layouts(prs)
    if ref is None or str(ref).strip() == "":
        raise ToolError("champ 'layout' manquant")
    sh = prs.slide_height
    ref = str(ref).strip()
    chosen = None
    if ref.startswith("@"):
        role = ref[1:].lower()
        alias = cfg.get("layout_aliases", {}).get(role)
        if alias:
            return resolve_layout(prs, alias, dict(cfg, layout_aliases={}))
        for lr in layouts:
            if classify_layout(lr.layout, sh) == role:
                chosen = lr
                break
        if chosen is None:
            roles = sorted(set(classify_layout(lr.layout, sh) for lr in layouts))
            raise ToolError("aucun layout pour le role '@%s' (roles detectes : %s). "
                            "Utilisez le nom exact du catalogue ou definissez 'layout_aliases' dans config/charte.json"
                            % (role, ", ".join(roles)))
    elif re.fullmatch(r"#?\d+", ref):
        i = int(ref.lstrip("#"))
        for lr in layouts:
            if lr.index == i:
                chosen = lr
        if chosen is None:
            raise ToolError("layout #%d inexistant (1..%d)" % (i, len(layouts)))
    else:
        exact = [lr for lr in layouts if lr.name == ref]
        folded = [lr for lr in layouts if _fold(lr.name) == _fold(ref)]
        if exact:
            chosen = exact[0]
        elif folded:
            chosen = folded[0]
        else:
            names = [lr.name for lr in layouts]
            close = difflib.get_close_matches(ref, names, n=3, cutoff=0.4)
            hint = (" Vouliez-vous dire : %s ?" % ", ".join('"%s"' % c for c in close)) if close else ""
            raise ToolError('layout "%s" absent du template.%s Voir assets/template-catalog.md' % (ref, hint))
    allowed = cfg.get("allowed_layouts")
    if allowed and chosen.name not in allowed:
        raise ToolError('layout "%s" non autorise par la charte (autorises : %s)' % (chosen.name, ", ".join(allowed)))
    if chosen.name in (cfg.get("forbidden_layouts") or []):
        raise ToolError('layout "%s" interdit par la charte' % chosen.name)
    return chosen


# ---------------------------------------------------------------- misc

def cm(v):
    return int(round(float(v) * EMU_PER_CM))


def to_cm(emu):
    return round((emu or 0) / EMU_PER_CM, 2)


def drop_all_slides(prs):
    """Supprime les diapositives d'exemple du modele (et sections/custom shows)."""
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)
    root = prs.part._element
    for ext in root.xpath('.//p:extLst/p:ext[@uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"]'):
        ext.getparent().remove(ext)
    for el in root.xpath('./p:custShowLst'):
        root.remove(el)


def resolve_path(p, *bases):
    """Chemin utilisateur -> Path existant (relatif a plusieurs bases), sinon None."""
    if p is None:
        return None
    pp = Path(str(p)).expanduser()
    if pp.is_absolute():
        return pp if pp.exists() else None
    for b in bases:
        if b is None:
            continue
        c = (Path(b) / pp)
        if c.exists():
            return c.resolve()
    c = Path.cwd() / pp
    return c.resolve() if c.exists() else None
