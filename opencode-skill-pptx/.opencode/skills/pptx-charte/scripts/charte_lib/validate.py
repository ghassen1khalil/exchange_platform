# -*- coding: utf-8 -*-
"""Controle de conformite d'une presentation par rapport au template."""
from __future__ import annotations

import re
from lxml import etree
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER as PH

from .common import (NSMAP, collect_layouts, find_template, load_config, open_presentation, theme_info)
from . import pattern as pat
from .textfit import estimate_fill_ratio

ERR, WARN, INFO = "ERREUR", "AVERTISSEMENT", "INFO"


class Finding(object):
    def __init__(self, sev, slide, msg, shape=""):
        self.sev, self.slide, self.msg, self.shape = sev, slide, msg, shape

    def __str__(self):
        loc = "diapo %s" % self.slide if self.slide else "deck"
        if self.shape:
            loc += " / %s" % self.shape
        return "[%s] %s : %s" % (self.sev, loc, self.msg)

    def as_dict(self):
        return {"severity": self.sev, "slide": self.slide, "shape": self.shape, "message": self.msg}


def _theme_sig(t):
    return (tuple(sorted(t["colors"].items())), t["major_font"], t["minor_font"])


def validate(pptx_path, template_path=None, cfg=None):
    cfg = cfg or load_config()
    tpl_path = find_template(template_path, cfg)
    tpl = open_presentation(tpl_path)
    prs = open_presentation(pptx_path)
    F = []
    lim = cfg["limits"]

    # ---- coherence globale avec le template
    if (prs.slide_width, prs.slide_height) != (tpl.slide_width, tpl.slide_height):
        F.append(Finding(ERR, None, "format de diapo different du template"))
    tpl_theme, deck_theme = theme_info(tpl), theme_info(prs)
    if _theme_sig(tpl_theme) != _theme_sig(deck_theme):
        F.append(Finding(ERR, None, "theme (couleurs/polices) different du template : "
                                    "la presentation n'a pas ete creee depuis le template officiel"))
    tpl_layouts = {lr.name for lr in collect_layouts(tpl)}
    palette = set(tpl_theme["colors"].values()) | {c.upper().lstrip("#") for c in cfg.get("extra_palette", [])}
    theme_fonts = {tpl_theme["major_font"], tpl_theme["minor_font"]}
    # couleurs/polices utilisees par le design du template lui-meme (masters, layouts, diapos d'exemple) = charte
    tcols, tfonts = _design_tokens(tpl)
    palette |= tcols
    theme_fonts |= tfonts
    tol = 180000  # 0,5 cm : fonds « perdus » du template

    titles = {}
    for n, slide in enumerate(prs.slides, 1):
        layout = slide.slide_layout
        if layout.name not in tpl_layouts:
            F.append(Finding(ERR, n, "layout \"%s\" absent du template" % layout.name))

        words = 0
        has_title = False
        for shp in slide.shapes:
            x = shp.left or 0
            y = shp.top or 0
            w = shp.width or 0
            h = shp.height or 0
            if shp.left is not None and (x < -tol or y < -tol or x + w > prs.slide_width + tol or y + h > prs.slide_height + tol):
                F.append(Finding(ERR, n, "la forme depasse du cadre de la diapo", shp.name))

            if shp.is_placeholder:
                t = shp.placeholder_format.type
                if t in (PH.TITLE, PH.CENTER_TITLE) and shp.has_text_frame and shp.text_frame.text.strip():
                    has_title = True
                    key = shp.text_frame.text.strip().lower()
                    titles.setdefault(key, []).append(n)
                    if len(shp.text_frame.text) > lim["max_title_chars"]:
                        F.append(Finding(WARN, n, "titre tres long (%d car.)" % len(shp.text_frame.text), shp.name))
                if shp.has_text_frame and not shp.text_frame.text.strip() and \
                        t not in (PH.DATE, PH.FOOTER, PH.SLIDE_NUMBER):
                    F.append(Finding(WARN, n, "zone vide (a remplir ou supprimer)", shp.name))
            elif shp.shape_type not in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP):
                # forme libre ajoutee : se recouvre-t-elle avec un placeholder ?
                for other in slide.placeholders:
                    if other.width and _overlap(shp, other) > 0.15:
                        F.append(Finding(WARN, n, "chevauche la zone '%s'" % other.name, shp.name))
                        break

            if shp.has_text_frame:
                txt = shp.text_frame.text
                words += len(txt.split())
                low = txt.lower()
                for bad in cfg.get("forbidden_texts", []):
                    if bad in low:
                        F.append(Finding(ERR, n, "texte residuel du modele ('%s')" % bad, shp.name))
                if shp.is_placeholder and txt.strip():
                    r = estimate_fill_ratio(shp, lim["avg_char_width_em"], lim["line_height_em"])
                    if r > 1.08:
                        F.append(Finding(WARN, n, "debordement probable (~%d%% de la zone)" % int(r * 100), shp.name))
                    bullets = [p for p in shp.text_frame.paragraphs if (p.level or 0) == 0 and p.text.strip()]
                    if shp.placeholder_format.type in (PH.BODY, PH.OBJECT) and len(bullets) > lim["max_bullets_per_placeholder"]:
                        F.append(Finding(WARN, n, "%d puces de niveau 1 (max %d)" % (
                            len(bullets), lim["max_bullets_per_placeholder"]), shp.name))

            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                descr = shp._element.xpath(".//p:cNvPr")[0].get("descr", "").strip()
                if not descr:
                    F.append(Finding(WARN, n, "image sans texte alternatif", shp.name))
                try:
                    iw, ih = shp.image.size
                    cl = (shp.crop_left or 0) + (shp.crop_right or 0)
                    ct = (shp.crop_top or 0) + (shp.crop_bottom or 0)
                    if w and h and ih and iw and abs((iw * (1 - cl)) / float(ih * (1 - ct)) - w / float(h)) > 0.06 * (w / float(h)):
                        F.append(Finding(WARN, n, "image deformee (ratio different de l'original)", shp.name))
                except Exception:
                    pass
            if getattr(shp, "has_chart", False) and shp.has_chart:
                if not shp._element.xpath(".//p:cNvPr")[0].get("descr", "").strip():
                    F.append(Finding(WARN, n, "graphique sans texte alternatif", shp.name))

        if not has_title:  # modele « diapo » : le titre est une zone de texte et non un placeholder
            try:
                t_slot, _ = pat.analyze_slide(slide, prs.slide_width, prs.slide_height)
            except Exception:
                t_slot = None
            if t_slot is not None and t_slot.sample.strip():
                has_title = True
                titles.setdefault(t_slot.sample.strip().lower(), []).append(n)
        # ---- mise en forme en dur
        xml = etree.tostring(slide._element).decode("utf-8", "ignore")
        for m in set(re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', _strip_graphics(slide))):
            if m.upper() not in palette:
                F.append(Finding(WARN, n, "couleur en dur #%s hors palette du theme" % m.upper()))
        for face in set(re.findall(r'<a:(?:latin|ea|cs) typeface="([^"]+)"', xml)):
            if not face.startswith("+") and face not in theme_fonts:
                F.append(Finding(WARN, n, "police en dur \"%s\" (theme : %s)" % (face, ", ".join(sorted(theme_fonts)))))
        minpt = cfg.get("min_font_pt", 12)
        for sz in set(int(v) for v in re.findall(r'\bsz="(\d+)"', xml)):
            if sz / 100.0 < minpt:
                F.append(Finding(WARN, n, "taille de police %.1f pt < minimum %s pt" % (sz / 100.0, minpt)))

        if words > lim["max_words_per_slide"]:
            F.append(Finding(WARN, n, "%d mots sur la diapo (max conseille %d) : trop dense" % (words, lim["max_words_per_slide"])))
        if not has_title:
            F.append(Finding(WARN, n, "pas de titre (navigation / accessibilite)"))

    for key, ns in titles.items():
        if len(ns) > 1:
            F.append(Finding(INFO, None, "titre identique sur les diapos %s : \"%s\"" % (ns, key[:50])))
    if not len(prs.slides):
        F.append(Finding(ERR, None, "aucune diapositive"))
    if not (prs.core_properties.title or "").strip():
        F.append(Finding(INFO, None, "titre du document (proprietes) vide"))
    return F, len(prs.slides)


def _design_tokens(tpl):
    """Couleurs hex et polices explicites presentes dans le template (donc conformes a la charte)."""
    cols, fonts = set(), set()
    parts = []
    for m in tpl.slide_masters:
        parts.append(m._element)
        parts.extend(l._element for l in m.slide_layouts)
    parts.extend(sl._element for sl in tpl.slides)
    for el in parts:
        xml = etree.tostring(el).decode("utf-8", "ignore")
        cols |= {c.upper() for c in re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', xml)}
        fonts |= {f for f in re.findall(r'<a:(?:latin|ea|cs) typeface="([^"]+)"', xml) if not f.startswith("+")}
    return cols, fonts


def _strip_graphics(slide):
    """XML de la diapo sans images/groupes (les icones du template ont leurs propres couleurs)."""
    import copy
    el = copy.deepcopy(slide._element)
    for node in el.xpath(".//p:pic | .//p:grpSp"):
        node.getparent().remove(node)
    return etree.tostring(el).decode("utf-8", "ignore")


def _overlap(a, b):
    ax1, ay1, ax2, ay2 = a.left or 0, a.top or 0, (a.left or 0) + (a.width or 0), (a.top or 0) + (a.height or 0)
    bx1, by1, bx2, by2 = b.left or 0, b.top or 0, (b.left or 0) + (b.width or 0), (b.top or 0) + (b.height or 0)
    iw, ih = min(ax2, bx2) - max(ax1, bx1), min(ay2, by2) - max(ay1, by1)
    if iw <= 0 or ih <= 0:
        return 0.0
    area = float((a.width or 1) * (a.height or 1))
    return iw * ih / area
