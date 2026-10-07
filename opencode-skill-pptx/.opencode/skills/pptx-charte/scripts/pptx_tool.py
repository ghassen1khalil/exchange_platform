# -*- coding: utf-8 -*-
"""Point d'entree unique du skill pptx-charte.

    run <commande> [options]       (run.cmd sous Windows, run.sh ailleurs)

Commandes : doctor | setup | catalog | icons | build | validate | render | outline

Ce fichier n'utilise que la bibliotheque standard tant que python-pptx n'est pas
disponible : il sait donc installer les dependances en espace utilisateur (sans droits admin).
"""
from __future__ import annotations

import argparse
import os
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL_DIR = HERE.parent
sys.path.insert(0, str(HERE))

REQUIREMENTS = ["python-pptx>=1.0.2", "lxml", "Pillow", "XlsxWriter", "typing_extensions"]


def libs_dir() -> Path:
    base = os.environ.get("LOCALAPPDATA") or os.environ.get("XDG_DATA_HOME") or str(Path.home() / ".local" / "share")
    tag = "py%d%d" % sys.version_info[:2]
    return Path(base) / "pptx-charte" / ("libs-" + tag)


def _activate_libs():
    d = libs_dir()
    if d.is_dir() and str(d) not in sys.path:
        sys.path.insert(0, str(d))


def _deps_ok() -> bool:
    try:
        import pptx  # noqa: F401
        import lxml  # noqa: F401
        import PIL  # noqa: F401
        return True
    except Exception:
        return False


_activate_libs()


def _utf8():
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


def need_deps():
    if not _deps_ok():
        print("[pptx-charte] Dependances Python manquantes (python-pptx, lxml, Pillow).\n"
              "  -> Lancez :  run setup\n"
              "     (installation en espace utilisateur, sans droits administrateur)", file=sys.stderr)
        sys.exit(3)


# ------------------------------------------------------------------ commandes

def cmd_setup(args):
    if sys.version_info < (3, 8):
        print("Python 3.8+ requis (version actuelle : %s)." % sys.version.split()[0], file=sys.stderr)
        return 2
    target = libs_dir()
    target.mkdir(parents=True, exist_ok=True)
    wheelhouse = args.wheelhouse or os.environ.get("PPTX_CHARTE_WHEELHOUSE")
    has_pip = subprocess.run([sys.executable, "-m", "pip", "--version"], capture_output=True).returncode == 0 and not os.environ.get("PPTX_CHARTE_NOPIP")
    if not has_pip:
        if not wheelhouse:
            print("[pptx-charte] pip est absent de ce Python (Python 'embarque' ?).\n"
                  "  -> fournissez un dossier de wheels : run setup --wheelhouse <dossier>\n"
                  "     (les .whl seront extraits sans pip ; voir references/windows-sans-admin.md)", file=sys.stderr)
            return 2
        import glob
        import zipfile
        wheels = sorted(glob.glob(os.path.join(wheelhouse, "*.whl")))
        if not wheels:
            print("[pptx-charte] Aucun .whl dans %s" % wheelhouse, file=sys.stderr)
            return 2
        print("pip absent : extraction directe de %d wheels dans %s" % (len(wheels), target))
        for w in wheels:
            with zipfile.ZipFile(w) as z:
                z.extractall(str(target))
        _activate_libs()
        print("\nVerification...")
        return cmd_doctor(args)
    cmd = [sys.executable, "-m", "pip", "install", "--upgrade", "--only-binary=:all:",
           "--target", str(target), "--disable-pip-version-check"]
    if wheelhouse:
        cmd += ["--no-index", "--find-links", wheelhouse]
    if args.index_url:
        cmd += ["--index-url", args.index_url]
    if args.proxy:
        cmd += ["--proxy", args.proxy]
    if args.trusted_host:
        for h in args.trusted_host:
            cmd += ["--trusted-host", h]
    cmd += REQUIREMENTS
    print("Installation dans : %s" % target)
    print("Commande : %s\n" % " ".join(cmd))
    r = subprocess.run(cmd)
    if r.returncode != 0:
        print("\n[pptx-charte] Echec de pip. Pistes (voir references/windows-sans-admin.md) :\n"
              "  * pas d'acces a Internet / PyPI bloque -> setup --wheelhouse <dossier de .whl> (installation hors-ligne)\n"
              "  * proxy d'entreprise -> setup --proxy http://proxy:port   ou variable HTTPS_PROXY\n"
              "  * miroir interne (Artifactory/Nexus) -> setup --index-url https://.../simple\n"
              "  * certificat d'entreprise -> setup --trusted-host pypi.org --trusted-host files.pythonhosted.org\n"
              "      ou: pip config set global.cert <chemin du .pem>", file=sys.stderr)
        return r.returncode
    _activate_libs()
    print("\nVerification...")
    return cmd_doctor(args)


def cmd_doctor(args):
    ok = True
    print("Python        : %s (%s)" % (sys.version.split()[0], sys.executable))
    print("Plateforme    : %s" % sys.platform)
    print("Dossier libs  : %s %s" % (libs_dir(), "(present)" if libs_dir().is_dir() else "(absent)"))
    if not _deps_ok():
        print("Dependances   : MANQUANTES -> lancez 'run setup'")
        return 3
    import lxml.etree
    import PIL
    import pptx
    print("python-pptx   : %s | lxml %s | Pillow %s" % (pptx.__version__, ".".join(map(str, lxml.etree.LXML_VERSION)), PIL.__version__))
    from charte_lib.common import ToolError, find_template, load_config, open_presentation
    try:
        cfg = load_config()
        print("Configuration : OK (config/charte.json)")
    except ToolError as e:
        print("Configuration : ERREUR - %s" % e)
        return 1
    try:
        tpl = find_template(getattr(args, 'template', None), cfg)
        prs = open_presentation(tpl)
        from charte_lib.common import collect_layouts
        print("Template      : %s (%d layouts, %d diapos d'exemple)" % (tpl.name, len(collect_layouts(prs)), len(prs.slides)))
    except ToolError as e:
        print("Template      : ABSENT/ILLISIBLE - %s" % e)
        ok = False
    cat = SKILL_DIR / "assets" / "template-catalog.md"
    print("Catalogue     : %s" % ("present" if cat.exists() else "absent -> lancez 'run catalog'"))
    out = Path.cwd() / "output"
    try:
        out.mkdir(exist_ok=True)
        t = out / ".w"
        t.write_text("x")
        t.unlink()
        print("Dossier output: %s (ecriture OK)" % out)
    except Exception as e:
        print("Dossier output: ECRITURE IMPOSSIBLE (%s)" % e)
        ok = False
    from charte_lib.render import available_engines
    eng = available_engines()
    print("Rendu apercu  : %s" % (", ".join(eng) if eng else "aucun moteur (PowerPoint/LibreOffice) -> 'validate' seulement"))
    return 0 if ok else 1


def cmd_catalog(args):
    need_deps()
    from charte_lib.catalog import inspect_template, to_markdown
    from charte_lib.common import find_template, load_config, open_presentation
    cfg = load_config()
    tpl = find_template(args.template, cfg)
    prs = open_presentation(tpl)
    info = inspect_template(prs, tpl, cfg)
    if args.render:
        from charte_lib.render import render
        try:
            _, files, sheet, outdir = render(tpl, SKILL_DIR / "assets" / "preview")
            info["preview"] = "assets/preview/" + Path(sheet).name if sheet else None
            print("Apercu du template : %d image(s) dans %s" % (len(files), outdir))
        except Exception as e:  # moteur absent : le catalogue reste valable
            print("[AVERTISSEMENT] Apercu impossible : %s" % e)
    md = to_markdown(info)
    if args.print:
        print(md)
        return 0
    out = Path(args.output) if args.output else SKILL_DIR / "assets" / "template-catalog.md"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(md, encoding="utf-8")
    print("Catalogue ecrit : %s  (%d layouts, %d icones)" % (out, len(info["layouts"]), len(info["icons"])))
    for lay in info["layouts"]:
        print("  %2d. %-38s -> @%s" % (lay["index"], '"%s"' % lay["name"], lay["role"]))
    return 0


def cmd_icons(args):
    need_deps()
    from charte_lib.common import find_template, load_config, open_presentation
    from charte_lib.icons import IconLibrary, describe
    cfg = load_config()
    prs = open_presentation(find_template(args.template, cfg))
    lib = IconLibrary.build(prs, cfg)
    items = lib.listing()
    if not items:
        print("Aucune icone. Ajoutez des PNG dans assets/icons/ ou reglez 'icon_library_slides' dans config/charte.json.")
        return 0
    for e in items:
        print(describe(e))
    return 0


def cmd_build(args):
    need_deps()
    from charte_lib.builder import build_from_spec
    from charte_lib.common import ToolError
    b, report, out, tpl = build_from_spec(args.spec, args.output, args.template, args.base)
    for w in report.warnings:
        print("[AVERTISSEMENT] %s" % w)
    if report.errors:
        print("\n%d ERREUR(S) - rien n'a ete ecrit. Corrigez la spec puis relancez :" % len(report.errors), file=sys.stderr)
        for e in report.errors:
            print("[ERREUR] %s" % e, file=sys.stderr)
        return 1
    if args.base and Path(args.base).resolve() == out.resolve():
        print("[ERREUR] La sortie ne peut pas etre le fichier --base : choisissez un autre nom (-o ...-v2.pptx).", file=sys.stderr)
        return 1
    b.save(out)
    print("\nPresentation creee : %s" % out.resolve())
    print("Template utilise   : %s" % Path(tpl).name)
    print("Diapositives (%d) :" % len(report.lines))
    for l in report.lines:
        print(l)
    print("\nEtape suivante : run validate \"%s\"" % out)
    return 0


def cmd_validate(args):
    need_deps()
    from charte_lib.validate import ERR, WARN, validate
    findings, n = validate(args.pptx, args.template)
    errs = [f for f in findings if f.sev == ERR]
    if args.json:
        import json
        print(json.dumps([f.as_dict() for f in findings], ensure_ascii=False, indent=2))
    else:
        order = {ERR: 0, WARN: 1}
        for f in sorted(findings, key=lambda f: (order.get(f.sev, 2), f.slide or 0)):
            print(f)
        print("\n%d diapositive(s) controlee(s) : %d erreur(s), %d avertissement(s)." % (
            n, len(errs), len([f for f in findings if f.sev == WARN])))
    return 1 if errs or (args.strict and any(f.sev == WARN for f in findings)) else 0


def cmd_render(args):
    need_deps()
    from charte_lib.render import render
    engine, files, sheet, out = render(args.pptx, args.out, args.engine)
    print("Moteur : %s" % engine)
    print("Images : %d dans %s" % (len(files), out))
    if sheet:
        print("Planche contact (toutes les diapos) : %s" % sheet)
    elif not files:
        print("Aucune image produite (PDF eventuellement dans %s)." % out)
    return 0


def cmd_outline(args):
    need_deps()
    from charte_lib.common import open_presentation
    prs = open_presentation(args.pptx)
    for n, s in enumerate(prs.slides, 1):
        title = ""
        if s.shapes.title is not None and s.shapes.title.has_text_frame:
            title = s.shapes.title.text_frame.text.replace("\n", " ")
        print("%2d. [%s] %s" % (n, s.slide_layout.name, title))
        for sh in s.placeholders:
            if sh.placeholder_format.type in (13, 15, 16):  # numero, pied de page, date
                continue
            if sh.has_text_frame and sh.text_frame.text.strip() and sh != s.shapes.title:
                print("      - (%s) %s" % (sh.placeholder_format.idx, sh.text_frame.text.strip().replace("\n", " | ")[:90]))
    return 0


def main(argv=None):
    _utf8()
    ap = argparse.ArgumentParser(prog="run", description="Skill pptx-charte : creation de PowerPoint conformes a la charte.")
    sub = ap.add_subparsers(dest="cmd")
    sub.required = True

    p = sub.add_parser("doctor", help="diagnostic de l'installation")
    p.add_argument("--template")
    p.set_defaults(func=cmd_doctor)

    p = sub.add_parser("setup", help="installe les dependances Python en espace utilisateur (sans admin)")
    p.add_argument("--wheelhouse", help="dossier de fichiers .whl (installation hors-ligne)")
    p.add_argument("--index-url", help="miroir PyPI interne")
    p.add_argument("--proxy", help="proxy http(s)://hote:port")
    p.add_argument("--trusted-host", action="append", help="hote a considerer comme sur (repetable)")
    p.set_defaults(func=cmd_setup)

    p = sub.add_parser("catalog", help="analyse le template et genere assets/template-catalog.md")
    p.add_argument("--template")
    p.add_argument("--output")
    p.add_argument("--print", action="store_true", help="affiche au lieu d'ecrire")
    p.add_argument("--render", action="store_true", help="genere aussi un apercu PNG des diapos du template (PowerPoint/LibreOffice)")
    p.set_defaults(func=cmd_catalog)

    p = sub.add_parser("icons", help="liste les icones utilisables")
    p.add_argument("--template")
    p.set_defaults(func=cmd_icons)

    p = sub.add_parser("build", help="genere un .pptx a partir d'une spec JSON")
    p.add_argument("spec")
    p.add_argument("-o", "--output")
    p.add_argument("--template")
    p.add_argument("--base", help="presentation existante a laquelle AJOUTER les diapos (conserve ses diapos)")
    p.set_defaults(func=cmd_build)

    p = sub.add_parser("validate", help="controle la conformite d'un .pptx")
    p.add_argument("pptx")
    p.add_argument("--template")
    p.add_argument("--json", action="store_true")
    p.add_argument("--strict", action="store_true", help="les avertissements font echouer")
    p.set_defaults(func=cmd_validate)

    p = sub.add_parser("render", help="apercu PNG (PowerPoint ou LibreOffice requis)")
    p.add_argument("pptx")
    p.add_argument("--out")
    p.add_argument("--engine", default="auto", choices=["auto", "powerpoint", "libreoffice"])
    p.set_defaults(func=cmd_render)

    p = sub.add_parser("outline", help="plan d'une presentation existante")
    p.add_argument("pptx")
    p.set_defaults(func=cmd_outline)

    args = ap.parse_args(argv)
    try:
        return args.func(args) or 0
    except Exception as e:
        name = type(e).__name__
        if name == "ToolError":
            print("[ERREUR] %s" % e, file=sys.stderr)
            return 1
        if os.environ.get("PPTX_CHARTE_DEBUG"):
            raise
        print("[ERREUR] Inattendue (%s : %s).\n  -> lancez 'run doctor' ; pour le detail technique : definir PPTX_CHARTE_DEBUG=1" % (name, e),
              file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
