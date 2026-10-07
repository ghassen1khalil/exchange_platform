# -*- coding: utf-8 -*-
"""Template de DEMO n°2 (fictif) dont le design est porte par des DIAPOS d'exemple en lorem ipsum
(formes posees a la main, mise en forme directe), sur le layout « Vierge » : cas d'un « modele » = simple .pptx.

    python make_demo_pattern_template.py [dossier_sortie]   ->   demo-pattern-template.pptx
"""
import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.util import Cm, Emu, Pt

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent
DARK, GREEN, ORANGE, GREY = RGBColor(0x12, 0x3B, 0x2E), RGBColor(0x00, 0x91, 0x5A), RGBColor(0xF2, 0xA9, 0x00), RGBColor(0x55, 0x5F, 0x5A)
FONT = "Georgia"
LOREM = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore."


def png(fn, size=256, bg=(0, 0, 0, 0)):
    im = Image.new("RGBA", (size, size), bg)
    fn(ImageDraw.Draw(im), size)
    b = io.BytesIO()
    im.save(b, "PNG")
    b.seek(0)
    return b


def ico(d, s):
    d.ellipse((20, 20, s - 20, s - 20), fill=(0, 145, 90, 255))
    d.rectangle((90, 90, 166, 166), fill=(255, 255, 255, 255))


def photo():
    im = Image.new("RGB", (1600, 1000), (222, 236, 229))
    d = ImageDraw.Draw(im)
    for i in range(0, 1600, 90):
        d.line((i, 0, i - 400, 1000), fill=(0, 145, 90), width=7)
    d.ellipse((560, 250, 1040, 730), fill=(18, 59, 46))
    b = io.BytesIO()
    im.save(b, "PNG")
    b.seek(0)
    return b


def textbox(slide, x, y, w, h, paragraphs, name):
    """paragraphs : [(texte, taille, gras, couleur, puce, niveau)]"""
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (text, size, bold, color, bullet, lvl) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = text
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = color
        if lvl:
            p.level = lvl
        if bullet:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", str(285750 + lvl * 285750))
            ppr.set("indent", "-285750")
            ppr.append(parse_xml('<a:buChar xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" char="&#8226;"/>'))
    return tb


def rect(slide, x, y, w, h, color, name):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    r.name = name
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()
    return r


def main():
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    blank = [l for l in prs.slide_layouts if l.name == "Blank"][0]
    W, H = 33.87, 19.05

    # 1. couverture
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, W, H, DARK, "Fond")
    rect(s, 2.5, 9.2, 4, 0.25, ORANGE, "Filet")
    s.shapes.add_picture(png(ico), Cm(28.5), Cm(1.5), Cm(3), Cm(3)).name = "Logo"
    textbox(s, 2.5, 5.0, 24, 4, [("Lorem ipsum dolor sit amet", 44, True, RGBColor(255, 255, 255), False, 0)], "Titre")
    textbox(s, 2.5, 9.8, 24, 2, [("Consectetur adipiscing elit - 2026", 20, False, RGBColor(0xCF, 0xE3, 0xD8), False, 0)], "Sous-titre")

    # 2. trois colonnes
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, W, 0.4, GREEN, "Bandeau")
    textbox(s, 2, 1.2, 29, 2.2, [("Lorem ipsum dolor sit amet consectetur", 32, True, DARK, False, 0)], "Titre")
    for i in range(3):
        x = 2 + i * 10.2
        s.shapes.add_picture(png(ico), Cm(x), Cm(4.4), Cm(2.2), Cm(2.2)).name = "Picto %d" % (i + 1)
        textbox(s, x + 2.6, 4.7, 2, 1.5, [("0%d" % (i + 1), 28, True, ORANGE, False, 0)], "Numero %d" % (i + 1))
        textbox(s, x, 7.2, 9, 8, [("Lorem ipsum dolor", 20, True, GREEN, False, 0), (LOREM, 14, False, GREY, False, 0)], "Colonne %d" % (i + 1))

    # 3. image + puces
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, W, 0.4, GREEN, "Bandeau")
    textbox(s, 2, 1.2, 29, 2.2, [("Lorem ipsum dolor sit amet", 32, True, DARK, False, 0)], "Titre")
    pic = s.shapes.add_picture(photo(), Cm(2), Cm(4.2), Cm(14), Cm(11.5))
    pic.name = "Visuel"
    textbox(s, 17.5, 4.2, 14.5, 11.5, [
        ("Lorem ipsum dolor sit amet consectetur", 18, False, DARK, True, 0),
        ("Sed do eiusmod tempor incididunt", 14, False, GREY, True, 1),
        ("Ut enim ad minim veniam quis nostrud", 18, False, DARK, True, 0),
        ("Duis aute irure dolor in reprehenderit", 18, False, DARK, True, 0)], "Puces")

    # 4. tableau
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, W, 0.4, GREEN, "Bandeau")
    textbox(s, 2, 1.2, 29, 2.2, [("Lorem ipsum dolor sit amet", 32, True, DARK, False, 0)], "Titre")
    gf = s.shapes.add_table(4, 3, Cm(2), Cm(4.5), Cm(29), Cm(6))
    for r in range(4):
        for c in range(3):
            gf.table.cell(r, c).text = "Lorem" if r else "Ipsum"

    # 5. graphique
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, W, 0.4, GREEN, "Bandeau")
    textbox(s, 2, 1.2, 29, 2.2, [("Lorem ipsum dolor sit amet", 32, True, DARK, False, 0)], "Titre")
    cd = CategoryChartData()
    cd.categories = ["A", "B", "C"]
    cd.add_series("Lorem", [1, 2, 3])
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Cm(2), Cm(4.2), Cm(29), Cm(12), cd)

    prs.core_properties.title = "Modele Acme (demo, design dans les diapos)"
    OUT.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT / "demo-pattern-template.pptx"))
    print("OK :", OUT / "demo-pattern-template.pptx")


if __name__ == "__main__":
    main()
