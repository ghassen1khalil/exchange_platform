# -*- coding: utf-8 -*-
"""Genere un template de DEMONSTRATION (fictif, "Acme Corp") pour tester le skill sans le vrai modele.

    python make_demo_template.py [dossier_sortie]

Produit demo-template.potx : 16:9, palette verte, police Arial, logo + bandeau sur le master,
une diapo d'exemple, une diapo "Bibliotheque d'icones" (3 images + 1 groupe vectoriel) et des sections.
"""
import io
import re
import sys
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.shapes.picture import CT_Picture
from pptx.oxml.shapes.autoshape import CT_Shape
from pptx.util import Emu, Pt

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent
SCALE = 12192000 / 9144000.0

PALETTE = {"dk2": "123B2E", "lt2": "EEF4F0", "accent1": "00915A", "accent2": "F2A900", "accent3": "1F6FB2",
           "accent4": "C8102E", "accent5": "6C7A89", "accent6": "7AC143", "hlink": "1F6FB2", "folHlink": "6C7A89"}


def png(draw_fn, size=256, bg=(0, 0, 0, 0)):
    im = Image.new("RGBA", (size, size), bg)
    draw_fn(ImageDraw.Draw(im), size)
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def logo(d, s):
    d.ellipse((8, 8, s - 8, s - 8), fill=(0, 145, 90, 255))
    d.rectangle((s * 0.3, s * 0.3, s * 0.7, s * 0.7), fill=(255, 255, 255, 255))


def icon_cloud(d, s):
    d.ellipse((30, 90, 130, 190), fill=(0, 145, 90, 255)); d.ellipse((90, 60, 200, 180), fill=(0, 145, 90, 255))
    d.rectangle((70, 130, 180, 190), fill=(0, 145, 90, 255))


def icon_lock(d, s):
    d.rectangle((60, 110, 196, 210), fill=(18, 59, 46, 255)); d.arc((85, 40, 171, 150), 180, 360, fill=(18, 59, 46, 255), width=14)


def icon_chart(d, s):
    for i, h in enumerate((60, 110, 160)):
        d.rectangle((50 + i * 55, 220 - h, 90 + i * 55, 220), fill=(242, 169, 0, 255))


def patch_theme(prs):
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    part = prs.slide_master.part.part_related_by(RT.THEME)
    xml = part.blob.decode("utf-8")
    for k, v in PALETTE.items():
        xml = re.sub(r"(<a:%s>)\s*<a:(?:srgbClr|sysClr)[^>]*/>\s*(</a:%s>)" % (k, k), r'\1<a:srgbClr val="%s"/>\2' % v, xml)
    xml = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*', r"\1Arial", xml)
    xml = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*', r"\1Arial", xml)
    xml = re.sub(r'<a:clrScheme name="[^"]*"', '<a:clrScheme name="Acme Corp"', xml)
    part._blob = xml.encode("utf-8")


def rescale_widescreen(prs):
    prs.slide_width = Emu(12192000)
    for holder in [prs.slide_master] + list(prs.slide_layouts):
        for sh in holder.shapes:
            own = sh._element.xpath("./p:spPr/a:xfrm") or sh._element.xpath("./p:xfrm")
            if own and sh.left is not None and sh.width is not None:
                l, t, w, h = sh.left, sh.top, sh.width, sh.height
                sh.left, sh.top, sh.width, sh.height = Emu(int(l * SCALE)), Emu(t), Emu(int(w * SCALE)), Emu(h)


def add_master_decor(prs):
    master = prs.slide_master
    tree = master.shapes._spTree
    band = CT_Shape.new_autoshape_sp(900, "Bandeau haut", "rect", 0, 0, 12192000, 120000)
    band.xpath(".//p:style")[0].getparent().remove(band.xpath(".//p:style")[0])
    sp_pr = band.xpath("./p:spPr")[0]
    sp_pr.append(parse_xml('<a:solidFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:schemeClr val="accent1"/></a:solidFill>'))
    sp_pr.append(parse_xml('<a:ln xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:noFill/></a:ln>'))
    tree.insert(2, band)
    _, rId = master.part.get_or_add_image_part(io.BytesIO(png(logo, 256)))
    pic = CT_Picture.new_pic(901, "Logo Acme", "Logo Acme Corp", rId, 11300000, 230000, 560000, 560000)
    tree.append(pic)


def make_icon_slide(prs):
    layout = [l for l in prs.slide_layouts if l.name == "Title Only"][0]
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = "Bibliotheque d'icones"
    x = 900000
    for name, fn in (("icone-cloud", icon_cloud), ("icone-lock", icon_lock), ("Picture 3", icon_chart)):
        pic = s.shapes.add_picture(io.BytesIO(png(fn)), Emu(x), Emu(2200000), Emu(900000), Emu(900000))
        if name != "Picture 3":
            pic.name = name
            pic._element.xpath(".//p:cNvPr")[0].set("descr", name)
        else:  # icone anonyme, nommee par sa legende
            tb = s.shapes.add_textbox(Emu(x - 100000), Emu(3200000), Emu(1100000), Emu(300000))
            tb.text_frame.text = "Reporting"
        x += 2200000
    # icone vectorielle (groupe)
    grp = s.shapes.add_group_shape()
    a = grp.shapes.add_shape(MSO_SHAPE.OVAL, Emu(x), Emu(2200000), Emu(900000), Emu(900000))
    b = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x + 300000), Emu(2500000), Emu(300000), Emu(300000))
    for shp, scheme in ((a, "accent1"), (b, "lt1")):
        shp.fill.solid()
        shp.fill.fore_color.theme_color = {"accent1": 5, "lt1": 2}[scheme]
        shp.line.fill.background()
    grp.name = "icone-cible"
    grp._element.xpath(".//p:cNvPr")[0].set("descr", "icone-cible")


def make_sample_slide(prs):
    layout = [l for l in prs.slide_layouts if l.name == "Title and Content"][0]
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = "Exemple de contenu (a supprimer)"
    s.placeholders[1].text = "Texte d'exemple du modele"


def add_sections(prs):
    ids = [sld.get("id") for sld in prs.slides._sldIdLst]
    ext = ('<p:ext xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}">'
           '<p14:sectionLst xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">'
           '<p14:section name="Exemples" id="{11111111-2222-3333-4444-555555555555}"><p14:sldIdLst>%s</p14:sldIdLst></p14:section>'
           '</p14:sectionLst></p:ext>' % "".join('<p14:sldId id="%s"/>' % i for i in ids))
    root = prs.part._element
    lst = root.find("{http://schemas.openxmlformats.org/presentationml/2006/main}extLst")
    if lst is None:
        lst = parse_xml('<p:extLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>')
        root.append(lst)
    lst.append(parse_xml(ext))


def main():
    prs = Presentation()
    rescale_widescreen(prs)
    patch_theme(prs)
    add_master_decor(prs)
    make_sample_slide(prs)
    make_icon_slide(prs)
    add_sections(prs)
    prs.core_properties.title = "Template Acme Corp (demo)"
    prs.core_properties.author = "Acme Corp"
    buf = io.BytesIO()
    prs.save(buf)
    # -> .potx : on change le content-type du document principal
    ct_from = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
    ct_to = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    OUT.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(str(OUT / "demo-template.potx"), "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == "[Content_Types].xml":
                data = data.decode("utf-8").replace(ct_from, ct_to).encode("utf-8")
            zout.writestr(it, data)
    print("OK :", OUT / "demo-template.potx")


if __name__ == "__main__":
    main()
